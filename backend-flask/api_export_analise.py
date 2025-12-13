"""
API de Exportação de Análises Jurídicas
Gera relatórios em PDF e Word baseados em análises multi-agente
"""

import logging
from flask import request, jsonify, make_response
from datetime import datetime
import json
from io import BytesIO

logger = logging.getLogger(__name__)

def gerar_relatorio_pdf(dados_analise):
    """Gera relatório PDF da análise"""
    try:
        from reportlab.pdfgen import canvas
        from reportlab.lib.pagesizes import letter, A4
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors
        from reportlab.lib.units import inch
        
        # Criar buffer em memória
        buffer = BytesIO()
        
        # Configurar documento
        doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=1*inch)
        styles = getSampleStyleSheet()
        story = []
        
        # Estilo personalizado
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=18,
            textColor=colors.HexColor('#2c3e50'),
            alignment=1,  # Centro
            spaceAfter=30
        )
        
        subtitle_style = ParagraphStyle(
            'CustomSubtitle',
            parent=styles['Heading2'],
            fontSize=14,
            textColor=colors.HexColor('#34495e'),
            spaceBefore=20,
            spaceAfter=10
        )
        
        # Título principal
        titulo = Paragraph("RELATÓRIO DE ANÁLISE JURÍDICA MULTI-AGENTE", title_style)
        story.append(titulo)
        
        # Subtítulo
        subtitulo = Paragraph("Legal Design Pro V2 - Sistema Especializado", styles['Heading3'])
        story.append(subtitulo)
        story.append(Spacer(1, 20))
        
        # Metadados
        agora = datetime.now()
        metadados_data = [
            ['Data da Análise', agora.strftime('%d/%m/%Y às %H:%M')],
            ['Sistema', 'Legal Design Pro V2'],
            ['Tipo de Análise', dados_analise.get('tipo', 'Multi-Agente')],
            ['Status', 'Concluída ✓']
        ]
        
        metadados_table = Table(metadados_data, colWidths=[2*inch, 3*inch])
        metadados_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#ecf0f1')),
            ('TEXTCOLOR', (0, 0), (0, -1), colors.HexColor('#2c3e50')),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTNAME', (1, 0), (1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('GRID', (0, 0), (-1, -1), 1, colors.HexColor('#bdc3c7'))
        ]))
        story.append(metadados_table)
        story.append(Spacer(1, 30))
        
        # Resumo Executivo
        story.append(Paragraph("RESUMO EXECUTIVO", subtitle_style))
        resumo_texto = dados_analise.get('resumo_executivo', 'Análise jurídica realizada com sucesso utilizando sistema multi-agente.')
        story.append(Paragraph(resumo_texto, styles['Normal']))
        story.append(Spacer(1, 20))
        
        # Análises Detalhadas
        if 'analises' in dados_analise:
            story.append(Paragraph("ANÁLISES DETALHADAS", subtitle_style))
            
            for i, analise in enumerate(dados_analise['analises']):
                agente_titulo = f"Análise {i+1} - {analise.get('agente', 'Sistema IA')}"
                story.append(Paragraph(agente_titulo, styles['Heading4']))
                
                conteudo_analise = analise.get('resultado', analise.get('analise', 'Conteúdo da análise não disponível'))
                story.append(Paragraph(conteudo_analise, styles['Normal']))
                story.append(Spacer(1, 15))
        
        # Conclusões
        story.append(Paragraph("CONCLUSÕES E RECOMENDAÇÕES", subtitle_style))
        conclusoes = dados_analise.get('conclusoes', 'Análise concluída com sucesso. Consulte os detalhes acima para informações específicas.')
        story.append(Paragraph(conclusoes, styles['Normal']))
        
        # Rodapé
        story.append(Spacer(1, 30))
        rodape = Paragraph(
            f"Relatório gerado automaticamente em {agora.strftime('%d/%m/%Y às %H:%M')}<br/>Legal Design Pro V2 - Sistema Multi-Agente",
            styles['Normal']
        )
        story.append(rodape)
        
        # Construir PDF
        doc.build(story)
        buffer.seek(0)
        
        return buffer.getvalue()
        
    except Exception as e:
        logger.error(f"Erro ao gerar PDF: {e}")
        raise e

def gerar_relatorio_word(dados_analise):
    """Gera relatório Word da análise"""
    try:
        from docx import Document
        from docx.shared import Inches, Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        
        doc = Document()
        
        # Configurar margens
        section = doc.sections[0]
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        
        # Título principal
        titulo = doc.add_heading('RELATÓRIO DE ANÁLISE JURÍDICA MULTI-AGENTE', 0)
        titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Subtítulo
        subtitulo = doc.add_heading('Legal Design Pro V2 - Sistema Especializado', 2)
        subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph()  # Espaço
        
        # Metadados
        agora = datetime.now()
        metadados_table = doc.add_table(rows=4, cols=2)
        metadados_table.style = 'Table Grid'
        
        metadados_data = [
            ['Data da Análise', agora.strftime('%d/%m/%Y às %H:%M')],
            ['Sistema', 'Legal Design Pro V2'],
            ['Tipo de Análise', dados_analise.get('tipo', 'Multi-Agente')],
            ['Status', 'Concluída ✓']
        ]
        
        for i, (campo, valor) in enumerate(metadados_data):
            metadados_table.cell(i, 0).text = campo
            metadados_table.cell(i, 1).text = valor
            metadados_table.cell(i, 0).paragraphs[0].runs[0].bold = True
        
        doc.add_paragraph()  # Espaço
        
        # Resumo Executivo
        doc.add_heading('RESUMO EXECUTIVO', 1)
        resumo_texto = dados_analise.get('resumo_executivo', 'Análise jurídica realizada com sucesso utilizando sistema multi-agente.')
        doc.add_paragraph(resumo_texto)
        
        # Análises Detalhadas
        if 'analises' in dados_analise:
            doc.add_heading('ANÁLISES DETALHADAS', 1)
            
            for i, analise in enumerate(dados_analise['analises']):
                agente_titulo = f"Análise {i+1} - {analise.get('agente', 'Sistema IA')}"
                doc.add_heading(agente_titulo, 2)
                
                conteudo_analise = analise.get('resultado', analise.get('analise', 'Conteúdo da análise não disponível'))
                doc.add_paragraph(conteudo_analise)
        
        # Conclusões
        doc.add_heading('CONCLUSÕES E RECOMENDAÇÕES', 1)
        conclusoes = dados_analise.get('conclusoes', 'Análise concluída com sucesso. Consulte os detalhes acima para informações específicas.')
        doc.add_paragraph(conclusoes)
        
        # Rodapé
        doc.add_paragraph()
        footer_para = doc.add_paragraph()
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_para.add_run(f'Relatório gerado automaticamente em {agora.strftime("%d/%m/%Y às %H:%M")}\n').italic = True
        footer_para.add_run('Legal Design Pro V2 - Sistema Multi-Agente').italic = True
        
        # Salvar em BytesIO
        doc_io = BytesIO()
        doc.save(doc_io)
        doc_io.seek(0)
        
        return doc_io.getvalue()
        
    except Exception as e:
        logger.error(f"Erro ao gerar Word: {e}")
        raise e

def registrar_api_export(app):
    """Registra endpoints da API de exportação"""
    
    @app.route('/api/export/pdf', methods=['POST'])
    def exportar_analise_pdf():
        try:
            data = request.get_json()
            
            if not data:
                return jsonify({
                    'success': False,
                    'error': 'Dados não fornecidos'
                }), 400
            
            # Processar dados para formato esperado
            dados_processados = {
                'tipo': data.get('tipo', 'Multi-Agente'),
                'resumo_executivo': data.get('resumo_executivo', 'Análise jurídica realizada com sistema multi-agente.'),
                'analises': data.get('analises', data.get('resultados', [])),
                'conclusoes': data.get('conclusoes', 'Análise concluída com sucesso.')
            }
            
            # Gerar PDF
            pdf_content = gerar_relatorio_pdf(dados_processados)
            
            # Preparar resposta
            response = make_response(pdf_content)
            response.headers['Content-Type'] = 'application/pdf'
            response.headers['Content-Disposition'] = f'attachment; filename=analise_juridica_{datetime.now().strftime("%Y%m%d_%H%M")}.pdf'
            
            logger.info("✅ PDF gerado com sucesso")
            return response
            
        except Exception as e:
            logger.error(f"❌ Erro ao exportar PDF: {str(e)}")
            return jsonify({
                'success': False,
                'error': f'Erro ao gerar PDF: {str(e)}'
            }), 500
    
    @app.route('/api/export/word', methods=['POST'])
    def exportar_analise_word():
        try:
            data = request.get_json()
            
            if not data:
                return jsonify({
                    'success': False,
                    'error': 'Dados não fornecidos'
                }), 400
            
            # Processar dados para formato esperado
            dados_processados = {
                'tipo': data.get('tipo', 'Multi-Agente'),
                'resumo_executivo': data.get('resumo_executivo', 'Análise jurídica realizada com sistema multi-agente.'),
                'analises': data.get('analises', data.get('resultados', [])),
                'conclusoes': data.get('conclusoes', 'Análise concluída com sucesso.')
            }
            
            # Gerar Word
            word_content = gerar_relatorio_word(dados_processados)
            
            # Preparar resposta
            response = make_response(word_content)
            response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
            response.headers['Content-Disposition'] = f'attachment; filename=analise_juridica_{datetime.now().strftime("%Y%m%d_%H%M")}.docx'
            
            logger.info("✅ Word gerado com sucesso")
            return response
            
        except Exception as e:
            logger.error(f"❌ Erro ao exportar Word: {str(e)}")
            return jsonify({
                'success': False,
                'error': f'Erro ao gerar Word: {str(e)}'
            }), 500
    
    @app.route('/api/export/status', methods=['GET'])
    def status_export():
        """Status da API de exportação"""
        return jsonify({
            'status': 'ativo',
            'formatos_disponiveis': ['PDF', 'Word (DOCX)'],
            'bibliotecas': ['ReportLab', 'python-docx'],
            'versao': '1.0.0'
        })
    
    logger.info("✅ API de exportação registrada")