"""
Utilitário para exportação de relatórios de consenso em diferentes formatos
"""

import io
import json
from datetime import datetime
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch

class ExportadorRelatorioConsenso:
    """Classe para exportar relatórios de consenso em diferentes formatos"""
    
    def __init__(self):
        self.styles = getSampleStyleSheet()
        self._criar_estilos_personalizados()
    
    def _criar_estilos_personalizados(self):
        """Cria estilos personalizados para o PDF"""
        self.styles.add(ParagraphStyle(
            name='TituloCustom',
            parent=self.styles['Heading1'],
            fontSize=18,
            spaceAfter=30,
            textColor=colors.HexColor('#1e3a8a'),
            alignment=1  # Centro
        ))
        
        self.styles.add(ParagraphStyle(
            name='SubtituloCustom',
            parent=self.styles['Heading2'],
            fontSize=14,
            spaceAfter=12,
            textColor=colors.HexColor('#3b82f6')
        ))
        
        self.styles.add(ParagraphStyle(
            name='ConteudoCustom',
            parent=self.styles['Normal'],
            fontSize=11,
            spaceAfter=6,
            leftIndent=20
        ))

    def gerar_pdf(self, consenso):
        """Gera PDF do relatório de consenso"""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        story = []
        
        # Título
        titulo = Paragraph(f"Relatório de Consenso - {consenso.titulo_relatorio}", self.styles['TituloCustom'])
        story.append(titulo)
        story.append(Spacer(1, 12))
        
        # Informações básicas
        info_data = [
            ['ID do Relatório:', consenso.numero_relatorio],
            ['Análise Base:', consenso.analise_numero_registro],
            ['Data de Geração:', consenso.data_criacao.strftime('%d/%m/%Y %H:%M:%S') if consenso.data_criacao else 'N/A'],
            ['Tempo de Processamento:', f"{consenso.tempo_processamento:.2f}s" if consenso.tempo_processamento else 'N/A'],
            ['Nível de Concordância:', f"{consenso.nivel_concordancia:.1f}%" if consenso.nivel_concordancia else 'N/A']
        ]
        
        info_table = Table(info_data, colWidths=[2*inch, 3*inch])
        info_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#f8f9fa')),
            ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('GRID', (0, 0), (-1, -1), 1, colors.black)
        ]))
        
        story.append(info_table)
        story.append(Spacer(1, 20))
        
        # Consenso Geral
        if consenso.consenso_geral:
            story.append(Paragraph("Consenso Geral", self.styles['SubtituloCustom']))
            story.append(Paragraph(consenso.consenso_geral, self.styles['ConteudoCustom']))
            story.append(Spacer(1, 15))
        
        # Pontos de Convergência
        if consenso.pontos_convergencia:
            story.append(Paragraph("Pontos de Convergência", self.styles['SubtituloCustom']))
            for ponto in consenso.pontos_convergencia:
                story.append(Paragraph(f"• {ponto}", self.styles['ConteudoCustom']))
            story.append(Spacer(1, 15))
        
        # Pontos de Divergência
        if consenso.pontos_divergencia:
            story.append(Paragraph("Pontos de Divergência", self.styles['SubtituloCustom']))
            for ponto in consenso.pontos_divergencia:
                story.append(Paragraph(f"• {ponto}", self.styles['ConteudoCustom']))
            story.append(Spacer(1, 15))
        
        # Análise de Riscos
        self._adicionar_secao_riscos(story, consenso)
        
        # Melhorias Recomendadas
        self._adicionar_secao_melhorias(story, consenso)
        
        # Estratégias por Prazo
        self._adicionar_secao_estrategias(story, consenso)
        
        doc.build(story)
        buffer.seek(0)
        return buffer.getvalue()
    
    def _adicionar_secao_riscos(self, story, consenso):
        """Adiciona seção de riscos ao PDF"""
        story.append(Paragraph("Análise de Riscos", self.styles['SubtituloCustom']))
        
        if consenso.riscos_criticos:
            story.append(Paragraph("Riscos Críticos:", self.styles['Heading3']))
            for risco in consenso.riscos_criticos:
                texto = risco if isinstance(risco, str) else risco.get('risco', str(risco))
                story.append(Paragraph(f"• {texto}", self.styles['ConteudoCustom']))
        
        if consenso.riscos_moderados:
            story.append(Paragraph("Riscos Moderados:", self.styles['Heading3']))
            for risco in consenso.riscos_moderados:
                texto = risco if isinstance(risco, str) else risco.get('risco', str(risco))
                story.append(Paragraph(f"• {texto}", self.styles['ConteudoCustom']))
        
        if consenso.riscos_baixos:
            story.append(Paragraph("Riscos Baixos:", self.styles['Heading3']))
            for risco in consenso.riscos_baixos:
                texto = risco if isinstance(risco, str) else risco.get('risco', str(risco))
                story.append(Paragraph(f"• {texto}", self.styles['ConteudoCustom']))
        
        story.append(Spacer(1, 15))
    
    def _adicionar_secao_melhorias(self, story, consenso):
        """Adiciona seção de melhorias ao PDF"""
        story.append(Paragraph("Melhorias Recomendadas", self.styles['SubtituloCustom']))
        
        if consenso.melhorias_urgentes:
            story.append(Paragraph("Urgentes:", self.styles['Heading3']))
            for melhoria in consenso.melhorias_urgentes:
                texto = melhoria if isinstance(melhoria, str) else melhoria.get('melhoria', str(melhoria))
                story.append(Paragraph(f"• {texto}", self.styles['ConteudoCustom']))
        
        if consenso.melhorias_importantes:
            story.append(Paragraph("Importantes:", self.styles['Heading3']))
            for melhoria in consenso.melhorias_importantes:
                texto = melhoria if isinstance(melhoria, str) else melhoria.get('melhoria', str(melhoria))
                story.append(Paragraph(f"• {texto}", self.styles['ConteudoCustom']))
        
        if consenso.melhorias_sugeridas:
            story.append(Paragraph("Sugeridas:", self.styles['Heading3']))
            for melhoria in consenso.melhorias_sugeridas:
                texto = melhoria if isinstance(melhoria, str) else melhoria.get('melhoria', str(melhoria))
                story.append(Paragraph(f"• {texto}", self.styles['ConteudoCustom']))
        
        story.append(Spacer(1, 15))
    
    def _adicionar_secao_estrategias(self, story, consenso):
        """Adiciona seção de estratégias ao PDF"""
        story.append(Paragraph("Estratégias Recomendadas", self.styles['SubtituloCustom']))
        
        if consenso.estrategias_curto_prazo:
            story.append(Paragraph("Curto Prazo (0-3 meses):", self.styles['Heading3']))
            for estrategia in consenso.estrategias_curto_prazo:
                texto = estrategia if isinstance(estrategia, str) else estrategia.get('estrategia', str(estrategia))
                story.append(Paragraph(f"• {texto}", self.styles['ConteudoCustom']))
        
        if consenso.estrategias_medio_prazo:
            story.append(Paragraph("Médio Prazo (3-12 meses):", self.styles['Heading3']))
            for estrategia in consenso.estrategias_medio_prazo:
                texto = estrategia if isinstance(estrategia, str) else estrategia.get('estrategia', str(estrategia))
                story.append(Paragraph(f"• {texto}", self.styles['ConteudoCustom']))
        
        if consenso.estrategias_longo_prazo:
            story.append(Paragraph("Longo Prazo (12+ meses):", self.styles['Heading3']))
            for estrategia in consenso.estrategias_longo_prazo:
                texto = estrategia if isinstance(estrategia, str) else estrategia.get('estrategia', str(estrategia))
                story.append(Paragraph(f"• {texto}", self.styles['ConteudoCustom']))

    def gerar_docx(self, consenso):
        """Gera DOCX do relatório de consenso"""
        document = Document()
        
        # Título
        title = document.add_heading(f'Relatório de Consenso - {consenso.titulo_relatorio}', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Informações básicas
        document.add_heading('Informações do Relatório', level=1)
        table = document.add_table(rows=5, cols=2)
        table.style = 'Table Grid'
        
        info_data = [
            ('ID do Relatório:', consenso.numero_relatorio),
            ('Análise Base:', consenso.analise_numero_registro),
            ('Data de Geração:', consenso.data_criacao.strftime('%d/%m/%Y %H:%M:%S') if consenso.data_criacao else 'N/A'),
            ('Tempo de Processamento:', f"{consenso.tempo_processamento:.2f}s" if consenso.tempo_processamento else 'N/A'),
            ('Nível de Concordância:', f"{consenso.nivel_concordancia:.1f}%" if consenso.nivel_concordancia else 'N/A')
        ]
        
        for i, (label, value) in enumerate(info_data):
            table.cell(i, 0).text = label
            table.cell(i, 1).text = str(value)
        
        # Consenso Geral
        if consenso.consenso_geral:
            document.add_heading('Consenso Geral', level=1)
            document.add_paragraph(consenso.consenso_geral)
        
        # Pontos de Convergência
        if consenso.pontos_convergencia:
            document.add_heading('Pontos de Convergência', level=1)
            for ponto in consenso.pontos_convergencia:
                p = document.add_paragraph()
                p.add_run('• ').bold = True
                p.add_run(ponto)
        
        # Pontos de Divergência
        if consenso.pontos_divergencia:
            document.add_heading('Pontos de Divergência', level=1)
            for ponto in consenso.pontos_divergencia:
                p = document.add_paragraph()
                p.add_run('• ').bold = True
                p.add_run(ponto)
        
        # Análise de Riscos
        self._adicionar_secao_riscos_docx(document, consenso)
        
        # Melhorias Recomendadas
        self._adicionar_secao_melhorias_docx(document, consenso)
        
        # Estratégias Recomendadas
        self._adicionar_secao_estrategias_docx(document, consenso)
        
        buffer = io.BytesIO()
        document.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
    
    def _adicionar_secao_riscos_docx(self, document, consenso):
        """Adiciona seção de riscos ao DOCX"""
        document.add_heading('Análise de Riscos', level=1)
        
        if consenso.riscos_criticos:
            document.add_heading('Riscos Críticos', level=2)
            for risco in consenso.riscos_criticos:
                texto = risco if isinstance(risco, str) else risco.get('risco', str(risco))
                p = document.add_paragraph()
                p.add_run('• ').bold = True
                p.add_run(texto)
        
        if consenso.riscos_moderados:
            document.add_heading('Riscos Moderados', level=2)
            for risco in consenso.riscos_moderados:
                texto = risco if isinstance(risco, str) else risco.get('risco', str(risco))
                p = document.add_paragraph()
                p.add_run('• ').bold = True
                p.add_run(texto)
        
        if consenso.riscos_baixos:
            document.add_heading('Riscos Baixos', level=2)
            for risco in consenso.riscos_baixos:
                texto = risco if isinstance(risco, str) else risco.get('risco', str(risco))
                p = document.add_paragraph()
                p.add_run('• ').bold = True
                p.add_run(texto)
    
    def _adicionar_secao_melhorias_docx(self, document, consenso):
        """Adiciona seção de melhorias ao DOCX"""
        document.add_heading('Melhorias Recomendadas', level=1)
        
        if consenso.melhorias_urgentes:
            document.add_heading('Urgentes', level=2)
            for melhoria in consenso.melhorias_urgentes:
                texto = melhoria if isinstance(melhoria, str) else melhoria.get('melhoria', str(melhoria))
                p = document.add_paragraph()
                p.add_run('• ').bold = True
                p.add_run(texto)
        
        if consenso.melhorias_importantes:
            document.add_heading('Importantes', level=2)
            for melhoria in consenso.melhorias_importantes:
                texto = melhoria if isinstance(melhoria, str) else melhoria.get('melhoria', str(melhoria))
                p = document.add_paragraph()
                p.add_run('• ').bold = True
                p.add_run(texto)
        
        if consenso.melhorias_sugeridas:
            document.add_heading('Sugeridas', level=2)
            for melhoria in consenso.melhorias_sugeridas:
                texto = melhoria if isinstance(melhoria, str) else melhoria.get('melhoria', str(melhoria))
                p = document.add_paragraph()
                p.add_run('• ').bold = True
                p.add_run(texto)
    
    def _adicionar_secao_estrategias_docx(self, document, consenso):
        """Adiciona seção de estratégias ao DOCX"""
        document.add_heading('Estratégias Recomendadas', level=1)
        
        if consenso.estrategias_curto_prazo:
            document.add_heading('Curto Prazo (0-3 meses)', level=2)
            for estrategia in consenso.estrategias_curto_prazo:
                texto = estrategia if isinstance(estrategia, str) else estrategia.get('estrategia', str(estrategia))
                p = document.add_paragraph()
                p.add_run('• ').bold = True
                p.add_run(texto)
        
        if consenso.estrategias_medio_prazo:
            document.add_heading('Médio Prazo (3-12 meses)', level=2)
            for estrategia in consenso.estrategias_medio_prazo:
                texto = estrategia if isinstance(estrategia, str) else estrategia.get('estrategia', str(estrategia))
                p = document.add_paragraph()
                p.add_run('• ').bold = True
                p.add_run(texto)
        
        if consenso.estrategias_longo_prazo:
            document.add_heading('Longo Prazo (12+ meses)', level=2)
            for estrategia in consenso.estrategias_longo_prazo:
                texto = estrategia if isinstance(estrategia, str) else estrategia.get('estrategia', str(estrategia))
                p = document.add_paragraph()
                p.add_run('• ').bold = True
                p.add_run(texto)

    def gerar_txt(self, consenso):
        """Gera TXT do relatório de consenso"""
        linhas = []
        linhas.append("="*80)
        linhas.append(f"RELATÓRIO DE CONSENSO - {consenso.titulo_relatorio}")
        linhas.append("="*80)
        linhas.append("")
        
        # Informações básicas
        linhas.append("INFORMAÇÕES DO RELATÓRIO")
        linhas.append("-"*30)
        linhas.append(f"ID do Relatório: {consenso.numero_relatorio}")
        linhas.append(f"Análise Base: {consenso.analise_numero_registro}")
        linhas.append(f"Data de Geração: {consenso.data_criacao.strftime('%d/%m/%Y %H:%M:%S') if consenso.data_criacao else 'N/A'}")
        linhas.append(f"Tempo de Processamento: {consenso.tempo_processamento:.2f}s" if consenso.tempo_processamento else 'N/A')
        linhas.append(f"Nível de Concordância: {consenso.nivel_concordancia:.1f}%" if consenso.nivel_concordancia else 'N/A')
        linhas.append("")
        
        # Consenso Geral
        if consenso.consenso_geral:
            linhas.append("CONSENSO GERAL")
            linhas.append("-"*15)
            linhas.append(consenso.consenso_geral)
            linhas.append("")
        
        # Pontos de Convergência
        if consenso.pontos_convergencia:
            linhas.append("PONTOS DE CONVERGÊNCIA")
            linhas.append("-"*25)
            for ponto in consenso.pontos_convergencia:
                linhas.append(f"• {ponto}")
            linhas.append("")
        
        # Pontos de Divergência
        if consenso.pontos_divergencia:
            linhas.append("PONTOS DE DIVERGÊNCIA")
            linhas.append("-"*23)
            for ponto in consenso.pontos_divergencia:
                linhas.append(f"• {ponto}")
            linhas.append("")
        
        # Análise de Riscos
        self._adicionar_secao_riscos_txt(linhas, consenso)
        
        # Melhorias Recomendadas
        self._adicionar_secao_melhorias_txt(linhas, consenso)
        
        # Estratégias Recomendadas
        self._adicionar_secao_estrategias_txt(linhas, consenso)
        
        return "\n".join(linhas)
    
    def _adicionar_secao_riscos_txt(self, linhas, consenso):
        """Adiciona seção de riscos ao TXT"""
        linhas.append("ANÁLISE DE RISCOS")
        linhas.append("-"*18)
        
        if consenso.riscos_criticos:
            linhas.append("\nRiscos Críticos:")
            for risco in consenso.riscos_criticos:
                texto = risco if isinstance(risco, str) else risco.get('risco', str(risco))
                linhas.append(f"• {texto}")
        
        if consenso.riscos_moderados:
            linhas.append("\nRiscos Moderados:")
            for risco in consenso.riscos_moderados:
                texto = risco if isinstance(risco, str) else risco.get('risco', str(risco))
                linhas.append(f"• {texto}")
        
        if consenso.riscos_baixos:
            linhas.append("\nRiscos Baixos:")
            for risco in consenso.riscos_baixos:
                texto = risco if isinstance(risco, str) else risco.get('risco', str(risco))
                linhas.append(f"• {texto}")
        
        linhas.append("")
    
    def _adicionar_secao_melhorias_txt(self, linhas, consenso):
        """Adiciona seção de melhorias ao TXT"""
        linhas.append("MELHORIAS RECOMENDADAS")
        linhas.append("-"*22)
        
        if consenso.melhorias_urgentes:
            linhas.append("\nUrgentes:")
            for melhoria in consenso.melhorias_urgentes:
                texto = melhoria if isinstance(melhoria, str) else melhoria.get('melhoria', str(melhoria))
                linhas.append(f"• {texto}")
        
        if consenso.melhorias_importantes:
            linhas.append("\nImportantes:")
            for melhoria in consenso.melhorias_importantes:
                texto = melhoria if isinstance(melhoria, str) else melhoria.get('melhoria', str(melhoria))
                linhas.append(f"• {texto}")
        
        if consenso.melhorias_sugeridas:
            linhas.append("\nSugeridas:")
            for melhoria in consenso.melhorias_sugeridas:
                texto = melhoria if isinstance(melhoria, str) else melhoria.get('melhoria', str(melhoria))
                linhas.append(f"• {texto}")
        
        linhas.append("")
    
    def _adicionar_secao_estrategias_txt(self, linhas, consenso):
        """Adiciona seção de estratégias ao TXT"""
        linhas.append("ESTRATÉGIAS RECOMENDADAS")
        linhas.append("-"*26)
        
        if consenso.estrategias_curto_prazo:
            linhas.append("\nCurto Prazo (0-3 meses):")
            for estrategia in consenso.estrategias_curto_prazo:
                texto = estrategia if isinstance(estrategia, str) else estrategia.get('estrategia', str(estrategia))
                linhas.append(f"• {texto}")
        
        if consenso.estrategias_medio_prazo:
            linhas.append("\nMédio Prazo (3-12 meses):")
            for estrategia in consenso.estrategias_medio_prazo:
                texto = estrategia if isinstance(estrategia, str) else estrategia.get('estrategia', str(estrategia))
                linhas.append(f"• {texto}")
        
        if consenso.estrategias_longo_prazo:
            linhas.append("\nLongo Prazo (12+ meses):")
            for estrategia in consenso.estrategias_longo_prazo:
                texto = estrategia if isinstance(estrategia, str) else estrategia.get('estrategia', str(estrategia))
                linhas.append(f"• {texto}")