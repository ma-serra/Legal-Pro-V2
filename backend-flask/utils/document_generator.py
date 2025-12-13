"""
Gerador de documentos para os módulos jurídicos.
Suporte para exportação em PDF, Word (DOCX) e HTML.
"""
import os
import uuid
import logging
from datetime import datetime
from typing import Dict, Optional, List
from flask import current_app
from flask_login import current_user

# Importações para geração de documentos
try:
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
    from docx.enum.style import WD_STYLE_TYPE
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    logging.warning("python-docx não disponível - exportação DOCX desabilitada")

try:
    from weasyprint import HTML, CSS
    from weasyprint.text.fonts import FontConfiguration
    WEASYPRINT_AVAILABLE = True
except ImportError:
    WEASYPRINT_AVAILABLE = False
    logging.warning("WeasyPrint não disponível - exportação PDF via HTML desabilitada")

try:
    from reportlab.lib.pagesizes import A4, letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
    from reportlab.lib import colors
    from reportlab.lib.units import inch
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False
    logging.warning("ReportLab não disponível - exportação PDF direta desabilitada")

# from utils.conversa_storage import salvar_arquivo_modulo  # Evitar importação circular

logger = logging.getLogger(__name__)

class DocumentGenerator:
    """Gerador de documentos para módulos jurídicos."""
    
    def __init__(self):
        self.upload_dir = os.path.join(os.getcwd(), 'uploads', 'documentos_gerados')
        self._ensure_upload_dir()
    
    def _ensure_upload_dir(self):
        """Garante que o diretório de upload existe."""
        os.makedirs(self.upload_dir, exist_ok=True)
    
    def _generate_filename(self, base_name: str, extension: str) -> str:
        """Gera nome único para arquivo."""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        unique_id = str(uuid.uuid4())[:8]
        safe_name = "".join(c for c in base_name if c.isalnum() or c in (' ', '-', '_')).strip()
        return f"{safe_name}_{timestamp}_{unique_id}.{extension}"
    
    def generate_docx(self, titulo: str, conteudo: str, modulo_id: str,
                     template_origem: str = '', metadata: Dict = None) -> Optional[str]:
        """
        Gera documento Word (.docx).
        
        Args:
            titulo: Título do documento
            conteudo: Conteúdo em HTML ou texto
            modulo_id: ID do módulo jurídico
            template_origem: Template usado para gerar
            metadata: Metadados adicionais
        
        Returns:
            str: Caminho do arquivo gerado ou None se erro
        """
        if not DOCX_AVAILABLE:
            logger.error("python-docx não está disponível")
            return None
        
        try:
            # Criar documento
            doc = Document()
            
            # Configurar estilos
            styles = doc.styles
            
            # Estilo para título
            if 'Título Jurídico' not in [s.name for s in styles]:
                title_style = styles.add_style('Título Jurídico', WD_STYLE_TYPE.PARAGRAPH)
                title_font = title_style.font
                title_font.name = 'Arial'
                title_font.size = Pt(16)
                title_font.bold = True
                title_style.paragraph_format.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                title_style.paragraph_format.space_after = Pt(12)
            
            # Estilo para texto normal
            if 'Texto Jurídico' not in [s.name for s in styles]:
                normal_style = styles.add_style('Texto Jurídico', WD_STYLE_TYPE.PARAGRAPH)
                normal_font = normal_style.font
                normal_font.name = 'Times New Roman'
                normal_font.size = Pt(12)
                normal_style.paragraph_format.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
                normal_style.paragraph_format.line_spacing = 1.15
            
            # Adicionar título
            title_paragraph = doc.add_paragraph(titulo, style='Título Jurídico')
            
            # Adicionar informações do módulo
            info_paragraph = doc.add_paragraph()
            info_paragraph.add_run(f"Módulo: {modulo_id.title()}")
            info_paragraph.add_run(f"\nData: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
            if current_user.is_authenticated:
                info_paragraph.add_run(f"\nUsuário: {current_user.username}")
            if template_origem:
                info_paragraph.add_run(f"\nTemplate: {template_origem}")
            
            doc.add_paragraph()  # Espaço
            
            # Processar conteúdo (remover tags HTML básicas)
            import html2text
            h = html2text.HTML2Text()
            h.ignore_links = True
            h.ignore_images = True
            plain_content = h.handle(conteudo)
            
            # Adicionar conteúdo
            paragraphs = plain_content.split('\n\n')
            for para in paragraphs:
                if para.strip():
                    doc.add_paragraph(para.strip(), style='Texto Jurídico')
            
            # Adicionar rodapé
            footer_paragraph = doc.add_paragraph()
            footer_paragraph.add_run(f"\nDocumento gerado automaticamente pelo Sistema Jurídico")
            footer_paragraph.add_run(f"\nEm {datetime.now().strftime('%d/%m/%Y às %H:%M')}")
            
            # Salvar arquivo
            filename = self._generate_filename(titulo, 'docx')
            filepath = os.path.join(self.upload_dir, filename)
            doc.save(filepath)
            
            # Registrar no banco de dados
            if current_user.is_authenticated:
                salvar_arquivo_modulo(
                    user_id=current_user.id,
                    modulo_id=modulo_id,
                    nome_arquivo=filename,
                    tipo_arquivo='docx',
                    caminho_arquivo=filepath,
                    template_origem=template_origem,
                    conteudo_original=conteudo,
                    mime_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
                )
            
            logger.info(f"Documento DOCX gerado: {filename}")
            return filepath
            
        except Exception as e:
            logger.error(f"Erro ao gerar DOCX: {e}")
            return None
    
    def generate_pdf_weasyprint(self, titulo: str, conteudo: str, modulo_id: str,
                               template_origem: str = '', metadata: Dict = None) -> Optional[str]:
        """
        Gera PDF usando WeasyPrint (melhor para HTML complexo).
        
        Args:
            titulo: Título do documento
            conteudo: Conteúdo em HTML
            modulo_id: ID do módulo jurídico
            template_origem: Template usado para gerar
            metadata: Metadados adicionais
        
        Returns:
            str: Caminho do arquivo gerado ou None se erro
        """
        if not WEASYPRINT_AVAILABLE:
            logger.error("WeasyPrint não está disponível")
            return None
        
        try:
            # CSS para estilização jurídica
            css_styles = """
            @page {
                size: A4;
                margin: 2.5cm;
                @bottom-center {
                    content: "Página " counter(page) " de " counter(pages);
                    font-size: 10pt;
                    color: #666;
                }
            }
            
            body {
                font-family: 'Times New Roman', Times, serif;
                font-size: 12pt;
                line-height: 1.4;
                text-align: justify;
                color: #000;
            }
            
            h1 {
                font-size: 16pt;
                font-weight: bold;
                text-align: center;
                margin-bottom: 20pt;
                color: #000;
            }
            
            h2 {
                font-size: 14pt;
                font-weight: bold;
                margin-top: 15pt;
                margin-bottom: 10pt;
                color: #000;
            }
            
            .document-info {
                border-bottom: 1pt solid #ccc;
                padding-bottom: 10pt;
                margin-bottom: 20pt;
                font-size: 10pt;
                color: #666;
            }
            
            .footer {
                margin-top: 30pt;
                border-top: 1pt solid #ccc;
                padding-top: 10pt;
                font-size: 10pt;
                color: #666;
                text-align: center;
            }
            
            p {
                margin-bottom: 10pt;
                text-indent: 1.5cm;
            }
            
            ul, ol {
                margin-left: 2cm;
            }
            
            blockquote {
                margin: 15pt 2cm;
                font-style: italic;
                border-left: 3pt solid #ccc;
                padding-left: 10pt;
            }
            """
            
            # HTML completo
            html_content = f"""
            <!DOCTYPE html>
            <html lang="pt-BR">
            <head>
                <meta charset="UTF-8">
                <title>{titulo}</title>
            </head>
            <body>
                <h1>{titulo}</h1>
                
                <div class="document-info">
                    <strong>Módulo:</strong> {modulo_id.title()}<br>
                    <strong>Data:</strong> {datetime.now().strftime('%d/%m/%Y %H:%M')}<br>
                    {f'<strong>Usuário:</strong> {current_user.username}<br>' if current_user.is_authenticated else ''}
                    {f'<strong>Template:</strong> {template_origem}<br>' if template_origem else ''}
                </div>
                
                <div class="content">
                    {conteudo}
                </div>
                
                <div class="footer">
                    Documento gerado automaticamente pelo Sistema Jurídico<br>
                    Em {datetime.now().strftime('%d/%m/%Y às %H:%M')}
                </div>
            </body>
            </html>
            """
            
            # Gerar PDF
            filename = self._generate_filename(titulo, 'pdf')
            filepath = os.path.join(self.upload_dir, filename)
            
            font_config = FontConfiguration()
            html_doc = HTML(string=html_content)
            css_doc = CSS(string=css_styles, font_config=font_config)
            
            html_doc.write_pdf(filepath, stylesheets=[css_doc], font_config=font_config)
            
            # Registrar no banco de dados
            if current_user.is_authenticated:
                salvar_arquivo_modulo(
                    user_id=current_user.id,
                    modulo_id=modulo_id,
                    nome_arquivo=filename,
                    tipo_arquivo='pdf',
                    caminho_arquivo=filepath,
                    template_origem=template_origem,
                    conteudo_original=conteudo,
                    mime_type='application/pdf'
                )
            
            logger.info(f"PDF gerado com WeasyPrint: {filename}")
            return filepath
            
        except Exception as e:
            logger.error(f"Erro ao gerar PDF com WeasyPrint: {e}")
            return None
    
    def generate_pdf_reportlab(self, titulo: str, conteudo: str, modulo_id: str,
                              template_origem: str = '', metadata: Dict = None) -> Optional[str]:
        """
        Gera PDF usando ReportLab (para texto simples).
        
        Args:
            titulo: Título do documento
            conteudo: Conteúdo em texto ou HTML simples
            modulo_id: ID do módulo jurídico
            template_origem: Template usado para gerar
            metadata: Metadados adicionais
        
        Returns:
            str: Caminho do arquivo gerado ou None se erro
        """
        if not REPORTLAB_AVAILABLE:
            logger.error("ReportLab não está disponível")
            return None
        
        try:
            # Processar conteúdo HTML para texto
            import html2text
            h = html2text.HTML2Text()
            h.ignore_links = True
            h.ignore_images = True
            plain_content = h.handle(conteudo)
            
            # Criar documento
            filename = self._generate_filename(titulo, 'pdf')
            filepath = os.path.join(self.upload_dir, filename)
            
            doc = SimpleDocTemplate(filepath, pagesize=A4,
                                  rightMargin=72, leftMargin=72,
                                  topMargin=72, bottomMargin=18)
            
            # Estilos
            styles = getSampleStyleSheet()
            
            # Estilo personalizado para título
            title_style = ParagraphStyle(
                'CustomTitle',
                parent=styles['Heading1'],
                fontSize=16,
                spaceAfter=30,
                alignment=TA_CENTER,
                fontName='Helvetica-Bold'
            )
            
            # Estilo para texto jurídico
            juridico_style = ParagraphStyle(
                'Juridico',
                parent=styles['Normal'],
                fontSize=12,
                spaceAfter=12,
                alignment=TA_JUSTIFY,
                fontName='Times-Roman',
                leftIndent=36,
                firstLineIndent=36
            )
            
            # Construir história do documento
            story = []
            
            # Título
            story.append(Paragraph(titulo, title_style))
            story.append(Spacer(1, 12))
            
            # Informações do documento
            info_text = f"""
            <b>Módulo:</b> {modulo_id.title()}<br/>
            <b>Data:</b> {datetime.now().strftime('%d/%m/%Y %H:%M')}<br/>
            {f'<b>Usuário:</b> {current_user.username}<br/>' if current_user.is_authenticated else ''}
            {f'<b>Template:</b> {template_origem}<br/>' if template_origem else ''}
            """
            story.append(Paragraph(info_text, styles['Normal']))
            story.append(Spacer(1, 20))
            
            # Conteúdo
            paragraphs = plain_content.split('\n\n')
            for para in paragraphs:
                if para.strip():
                    story.append(Paragraph(para.strip(), juridico_style))
            
            # Rodapé
            story.append(Spacer(1, 30))
            footer_text = f"""
            <i>Documento gerado automaticamente pelo Sistema Jurídico<br/>
            Em {datetime.now().strftime('%d/%m/%Y às %H:%M')}</i>
            """
            story.append(Paragraph(footer_text, styles['Normal']))
            
            # Gerar PDF
            doc.build(story)
            
            # Registrar no banco de dados
            if current_user.is_authenticated:
                salvar_arquivo_modulo(
                    user_id=current_user.id,
                    modulo_id=modulo_id,
                    nome_arquivo=filename,
                    tipo_arquivo='pdf',
                    caminho_arquivo=filepath,
                    template_origem=template_origem,
                    conteudo_original=conteudo,
                    mime_type='application/pdf'
                )
            
            logger.info(f"PDF gerado com ReportLab: {filename}")
            return filepath
            
        except Exception as e:
            logger.error(f"Erro ao gerar PDF com ReportLab: {e}")
            return None
    
    def generate_html(self, titulo: str, conteudo: str, modulo_id: str,
                     template_origem: str = '', metadata: Dict = None) -> Optional[str]:
        """
        Gera arquivo HTML formatado.
        
        Args:
            titulo: Título do documento
            conteudo: Conteúdo em HTML
            modulo_id: ID do módulo jurídico
            template_origem: Template usado para gerar
            metadata: Metadados adicionais
        
        Returns:
            str: Caminho do arquivo gerado ou None se erro
        """
        try:
            # HTML completo com estilos
            html_content = f"""
            <!DOCTYPE html>
            <html lang="pt-BR">
            <head>
                <meta charset="UTF-8">
                <meta name="viewport" content="width=device-width, initial-scale=1.0">
                <title>{titulo}</title>
                <style>
                    body {{
                        font-family: 'Times New Roman', Times, serif;
                        max-width: 800px;
                        margin: 0 auto;
                        padding: 40px 20px;
                        line-height: 1.6;
                        color: #333;
                    }}
                    
                    h1 {{
                        text-align: center;
                        color: #2c3e50;
                        margin-bottom: 30px;
                        border-bottom: 2px solid #3498db;
                        padding-bottom: 10px;
                    }}
                    
                    .document-info {{
                        background: #f8f9fa;
                        border: 1px solid #dee2e6;
                        border-radius: 5px;
                        padding: 15px;
                        margin-bottom: 30px;
                        font-size: 14px;
                    }}
                    
                    .content {{
                        text-align: justify;
                        text-indent: 2em;
                    }}
                    
                    .footer {{
                        margin-top: 40px;
                        padding-top: 20px;
                        border-top: 1px solid #dee2e6;
                        text-align: center;
                        font-size: 12px;
                        color: #6c757d;
                    }}
                    
                    p {{
                        margin-bottom: 15px;
                    }}
                    
                    ul, ol {{
                        margin-left: 30px;
                    }}
                    
                    blockquote {{
                        border-left: 4px solid #3498db;
                        margin: 20px 0;
                        padding-left: 20px;
                        font-style: italic;
                    }}
                    
                    @media print {{
                        body {{
                            margin: 0;
                            padding: 20px;
                        }}
                        
                        .no-print {{
                            display: none;
                        }}
                    }}
                </style>
            </head>
            <body>
                <h1>{titulo}</h1>
                
                <div class="document-info">
                    <strong>Módulo:</strong> {modulo_id.title()}<br>
                    <strong>Data:</strong> {datetime.now().strftime('%d/%m/%Y %H:%M')}<br>
                    {f'<strong>Usuário:</strong> {current_user.username}<br>' if current_user.is_authenticated else ''}
                    {f'<strong>Template:</strong> {template_origem}<br>' if template_origem else ''}
                </div>
                
                <div class="content">
                    {conteudo}
                </div>
                
                <div class="footer">
                    <p>Documento gerado automaticamente pelo Sistema Jurídico</p>
                    <p>Em {datetime.now().strftime('%d/%m/%Y às %H:%M')}</p>
                </div>
            </body>
            </html>
            """
            
            # Salvar arquivo
            filename = self._generate_filename(titulo, 'html')
            filepath = os.path.join(self.upload_dir, filename)
            
            with open(filepath, 'w', encoding='utf-8') as f:
                f.write(html_content)
            
            # Registrar no banco de dados
            if current_user.is_authenticated:
                salvar_arquivo_modulo(
                    user_id=current_user.id,
                    modulo_id=modulo_id,
                    nome_arquivo=filename,
                    tipo_arquivo='html',
                    caminho_arquivo=filepath,
                    template_origem=template_origem,
                    conteudo_original=conteudo,
                    mime_type='text/html'
                )
            
            logger.info(f"HTML gerado: {filename}")
            return filepath
            
        except Exception as e:
            logger.error(f"Erro ao gerar HTML: {e}")
            return None
    
    def generate_document(self, titulo: str, conteudo: str, modulo_id: str,
                         formato: str = 'pdf', template_origem: str = '',
                         metadata: Dict = None) -> Optional[str]:
        """
        Gera documento no formato especificado.
        
        Args:
            titulo: Título do documento
            conteudo: Conteúdo em HTML
            modulo_id: ID do módulo jurídico
            formato: Formato desejado ('pdf', 'docx', 'html')
            template_origem: Template usado para gerar
            metadata: Metadados adicionais
        
        Returns:
            str: Caminho do arquivo gerado ou None se erro
        """
        formato = formato.lower()
        
        if formato == 'docx':
            return self.generate_docx(titulo, conteudo, modulo_id, template_origem, metadata)
        elif formato == 'pdf':
            # Tentar WeasyPrint primeiro, depois ReportLab
            filepath = self.generate_pdf_weasyprint(titulo, conteudo, modulo_id, template_origem, metadata)
            if not filepath:
                filepath = self.generate_pdf_reportlab(titulo, conteudo, modulo_id, template_origem, metadata)
            return filepath
        elif formato == 'html':
            return self.generate_html(titulo, conteudo, modulo_id, template_origem, metadata)
        else:
            logger.error(f"Formato não suportado: {formato}")
            return None
    
    def get_available_formats(self) -> List[str]:
        """Retorna lista de formatos disponíveis."""
        formats = ['html']  # HTML sempre disponível
        
        if DOCX_AVAILABLE:
            formats.append('docx')
        
        if WEASYPRINT_AVAILABLE or REPORTLAB_AVAILABLE:
            formats.append('pdf')
        
        return formats
    
    def gerar_pdf_transcricao(self, dados_transcricao: Dict) -> Optional[str]:
        """
        Gera PDF específico para transcrições com análise completa.
        
        Args:
            dados_transcricao: Dados completos da transcrição
        
        Returns:
            str: Caminho do arquivo gerado ou None se erro
        """
        try:
            # Extrair dados
            texto = dados_transcricao.get('text', '')
            segments = dados_transcricao.get('segments', [])
            sentiment = dados_transcricao.get('sentiment', {})
            resumo = sentiment.get('resumo_audio', '')
            principais_pontos = sentiment.get('principais_pontos', {})
            
            # Criar conteúdo HTML
            html_content = f"""
            <!DOCTYPE html>
            <html lang="pt-BR">
            <head>
                <meta charset="UTF-8">
                <title>Transcrição de Áudio - {datetime.now().strftime('%d/%m/%Y')}</title>
                <style>
                    body {{ font-family: Arial, sans-serif; margin: 40px; line-height: 1.6; }}
                    .header {{ text-align: center; margin-bottom: 30px; border-bottom: 2px solid #333; padding-bottom: 20px; }}
                    .section {{ margin-bottom: 30px; }}
                    .section h2 {{ color: #333; border-left: 4px solid #007bff; padding-left: 10px; }}
                    .segment {{ margin-bottom: 15px; padding: 10px; background-color: #f8f9fa; border-left: 3px solid #007bff; }}
                    .timestamp {{ font-weight: bold; color: #6c757d; }}
                    .resumo {{ background-color: #e9f7ff; padding: 15px; border-radius: 8px; }}
                    .pontos-principais {{ background-color: #fff3cd; padding: 15px; border-radius: 8px; }}
                    .sentimento {{ background-color: #f8f9fa; padding: 15px; border-radius: 8px; }}
                    .footer {{ margin-top: 50px; text-align: center; font-size: 0.9em; color: #6c757d; }}
                </style>
            </head>
            <body>
                <div class="header">
                    <h1>📝 Transcrição de Áudio Jurídico</h1>
                    <p><strong>Data de Processamento:</strong> {datetime.now().strftime('%d/%m/%Y às %H:%M')}</p>
                    <p><strong>Duração:</strong> {dados_transcricao.get('duration', 0):.1f} segundos</p>
                    <p><strong>Confiança:</strong> {dados_transcricao.get('confidence', 0):.1%}</p>
                </div>
            """
            
            # Adicionar resumo se disponível
            if resumo:
                html_content += f"""
                <div class="section">
                    <h2>📝 Resumo do Áudio</h2>
                    <div class="resumo">
                        <p>{resumo}</p>
                    </div>
                </div>
                """
            
            # Adicionar principais pontos se disponível
            if principais_pontos:
                html_content += f"""
                <div class="section">
                    <h2>🎯 Principais Pontos Identificados</h2>
                    <div class="pontos-principais">
                """
                
                if principais_pontos.get('frases_destacadas'):
                    html_content += "<h3>Frases Destacadas:</h3>"
                    for frase in principais_pontos['frases_destacadas']:
                        html_content += f'<p style="font-style: italic; color: #0056b3;">💬 "{frase}"</p>'
                
                if principais_pontos.get('pontos_principais'):
                    html_content += "<h3>Pontos Principais:</h3><ul>"
                    for ponto in principais_pontos['pontos_principais']:
                        html_content += f"<li>{ponto}</li>"
                    html_content += "</ul>"
                
                html_content += "</div></div>"
            
            # Adicionar análise de sentimento
            if sentiment:
                html_content += f"""
                <div class="section">
                    <h2>📊 Análise de Sentimento</h2>
                    <div class="sentimento">
                        <p><strong>Sentimento Principal:</strong> {sentiment.get('main_sentiment', 'N/A')}</p>
                """
                
                if sentiment.get('sentiment_percentages'):
                    perc = sentiment['sentiment_percentages']
                    html_content += f"""
                        <p><strong>Distribuição:</strong></p>
                        <ul>
                            <li>😊 Positivo: {perc.get('Positivo', 0)}%</li>
                            <li>😞 Negativo: {perc.get('Negativo', 0)}%</li>
                            <li>😐 Neutro: {perc.get('Neutro', 0)}%</li>
                        </ul>
                    """
                
                if sentiment.get('emotion_percentages'):
                    emoc = sentiment['emotion_percentages']
                    html_content += "<p><strong>Emoções Detectadas:</strong></p><ul>"
                    for emocao, perc in emoc.items():
                        html_content += f"<li>{emocao}: {perc}%</li>"
                    html_content += "</ul>"
                
                html_content += "</div></div>"
            
            # Adicionar transcrição completa
            html_content += f"""
                <div class="section">
                    <h2>🎤 Transcrição Completa</h2>
            """
            
            if segments:
                for segment in segments:
                    start_time = segment.get('start', 0)
                    end_time = segment.get('end', 0)
                    text = segment.get('text', '')
                    
                    start_min = int(start_time // 60)
                    start_sec = int(start_time % 60)
                    end_min = int(end_time // 60)
                    end_sec = int(end_time % 60)
                    
                    html_content += f"""
                    <div class="segment">
                        <span class="timestamp">[{start_min:02d}:{start_sec:02d} → {end_min:02d}:{end_sec:02d}]</span>
                        <p>{text}</p>
                    </div>
                    """
            else:
                html_content += f"<p>{texto}</p>"
            
            html_content += """
                </div>
                
                <div class="footer">
                    <p>Documento gerado automaticamente pelo Sistema Jurídico</p>
                    <p>Transcrição realizada com tecnologia Whisper AI</p>
                </div>
            </body>
            </html>
            """
            
            # Gerar PDF
            filename = f"transcricao_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
            filepath = os.path.join(self.upload_dir, filename)
            
            if WEASYPRINT_AVAILABLE:
                HTML(string=html_content).write_pdf(filepath)
            else:
                # Fallback para ReportLab
                return self._gerar_pdf_reportlab_transcricao(dados_transcricao)
            
            logger.info(f"PDF de transcrição gerado: {filename}")
            return filepath
            
        except Exception as e:
            logger.error(f"Erro ao gerar PDF de transcrição: {e}")
            return None
    
    def gerar_docx_transcricao(self, dados_transcricao: Dict) -> Optional[str]:
        """
        Gera DOCX específico para transcrições com análise completa.
        
        Args:
            dados_transcricao: Dados completos da transcrição
        
        Returns:
            str: Caminho do arquivo gerado ou None se erro
        """
        if not DOCX_AVAILABLE:
            logger.error("python-docx não está disponível")
            return None
        
        try:
            # Criar documento
            doc = Document()
            
            # Adicionar título
            title = doc.add_heading('📝 Transcrição de Áudio Jurídico', 0)
            title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            
            # Informações do documento  
            info_para = doc.add_paragraph()
            info_para.add_run(f"Data de Processamento: ").bold = True
            info_para.add_run(f"{datetime.now().strftime('%d/%m/%Y às %H:%M')}")
            info_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            
            # Adicionar espaçamento sem quebra de página
            doc.add_paragraph()
            doc.add_paragraph()
            
            # Extrair dados
            sentiment = dados_transcricao.get('sentiment', {})
            resumo = sentiment.get('resumo_audio', '')
            principais_pontos = sentiment.get('principais_pontos', {})
            segments = dados_transcricao.get('segments', [])
            texto = dados_transcricao.get('text', '')
            
            # Adicionar resumo se disponível
            if resumo:
                doc.add_heading('📝 Resumo do Áudio', level=1)
                para_resumo = doc.add_paragraph(resumo)
                para_resumo.style = 'Quote'
                doc.add_paragraph()
            
            # Adicionar principais pontos se disponível
            if principais_pontos:
                doc.add_heading('🎯 Principais Pontos Identificados', level=1)
                
                if principais_pontos.get('frases_destacadas'):
                    doc.add_heading('Frases Destacadas:', level=2)
                    for frase in principais_pontos['frases_destacadas']:
                        para_frase = doc.add_paragraph()
                        para_frase.add_run(f'💬 "{frase}"').italic = True
                
                if principais_pontos.get('pontos_principais'):
                    doc.add_heading('Pontos Principais:', level=2)
                    for ponto in principais_pontos['pontos_principais']:
                        doc.add_paragraph(ponto, style='List Bullet')
                
                doc.add_paragraph()
            
            # Adicionar análise de sentimento
            if sentiment:
                doc.add_heading('📊 Análise de Sentimento', level=1)
                
                sent_para = doc.add_paragraph()
                sent_para.add_run("Sentimento Principal: ").bold = True
                sent_para.add_run(sentiment.get('main_sentiment', 'N/A'))
                
                if sentiment.get('sentiment_percentages'):
                    perc = sentiment['sentiment_percentages']
                    doc.add_paragraph("Distribuição:", style='Heading 3')
                    doc.add_paragraph(f"😊 Positivo: {perc.get('Positivo', 0)}%", style='List Bullet')
                    doc.add_paragraph(f"😞 Negativo: {perc.get('Negativo', 0)}%", style='List Bullet')
                    doc.add_paragraph(f"😐 Neutro: {perc.get('Neutro', 0)}%", style='List Bullet')
                
                if sentiment.get('emotion_percentages'):
                    emoc = sentiment['emotion_percentages']
                    doc.add_paragraph("Emoções Detectadas:", style='Heading 3')
                    for emocao, perc in emoc.items():
                        doc.add_paragraph(f"{emocao}: {perc}%", style='List Bullet')
                
                doc.add_paragraph()
            
            # Adicionar transcrição completa
            doc.add_heading('🎤 Transcrição Completa', level=1)
            
            if segments:
                for segment in segments:
                    start_time = segment.get('start', 0)
                    end_time = segment.get('end', 0)
                    text = segment.get('text', '')
                    
                    start_min = int(start_time // 60)
                    start_sec = int(start_time % 60)
                    end_min = int(end_time // 60)
                    end_sec = int(end_time % 60)
                    
                    para_segment = doc.add_paragraph()
                    timestamp_run = para_segment.add_run(f"[{start_min:02d}:{start_sec:02d} → {end_min:02d}:{end_sec:02d}] ")
                    timestamp_run.bold = True
                    para_segment.add_run(text)
            else:
                doc.add_paragraph(texto)
            
            # Adicionar rodapé
            doc.add_paragraph()
            footer_para = doc.add_paragraph()
            footer_para.add_run("Documento gerado automaticamente pelo Sistema Jurídico").bold = True
            footer_para.add_run("Transcrição realizada com tecnologia Whisper AI")
            footer_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            
            # Salvar documento
            filename = f"transcricao_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
            filepath = os.path.join(self.upload_dir, filename)
            doc.save(filepath)
            
            logger.info(f"DOCX de transcrição gerado: {filename}")
            return filepath
            
        except Exception as e:
            logger.error(f"Erro ao gerar DOCX de transcrição: {e}")
            return None
    
    def _gerar_pdf_reportlab_transcricao(self, dados_transcricao: Dict) -> Optional[str]:
        """Fallback para gerar PDF com ReportLab quando WeasyPrint não está disponível."""
        if not REPORTLAB_AVAILABLE:
            return None
        
        try:
            filename = f"transcricao_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
            filepath = os.path.join(self.upload_dir, filename)
            
            doc = SimpleDocTemplate(filepath, pagesize=A4)
            styles = getSampleStyleSheet()
            story = []
            
            # Título
            title_style = ParagraphStyle(
                'CustomTitle',
                parent=styles['Heading1'],
                fontSize=18,
                alignment=TA_CENTER,
                spaceAfter=30
            )
            story.append(Paragraph("📝 Transcrição de Áudio Jurídico", title_style))
            
            # Informações
            info_text = f"""
            <b>Data de Processamento:</b> {datetime.now().strftime('%d/%m/%Y às %H:%M')}
            """
            story.append(Paragraph(info_text, styles['Normal']))
            story.append(Spacer(1, 20))
            
            # Transcrição
            sentiment = dados_transcricao.get('sentiment', {})
            if sentiment.get('resumo_audio'):
                story.append(Paragraph("<b>📝 Resumo:</b>", styles['Heading2']))
                story.append(Paragraph(sentiment['resumo_audio'], styles['Normal']))
                story.append(Spacer(1, 15))
            
            story.append(Paragraph("<b>🎤 Transcrição:</b>", styles['Heading2']))
            
            segments = dados_transcricao.get('segments', [])
            if segments:
                for segment in segments:
                    start_time = segment.get('start', 0)
                    end_time = segment.get('end', 0)
                    text = segment.get('text', '')
                    
                    start_min = int(start_time // 60)
                    start_sec = int(start_time % 60)
                    end_min = int(end_time // 60)
                    end_sec = int(end_time % 60)
                    
                    segment_text = f"<b>[{start_min:02d}:{start_sec:02d} → {end_min:02d}:{end_sec:02d}]</b> {text}"
                    story.append(Paragraph(segment_text, styles['Normal']))
            else:
                story.append(Paragraph(dados_transcricao.get('text', ''), styles['Normal']))
            
            doc.build(story)
            logger.info(f"PDF de transcrição gerado com ReportLab: {filename}")
            return filepath
            
        except Exception as e:
            logger.error(f"Erro ao gerar PDF com ReportLab: {e}")
            return None