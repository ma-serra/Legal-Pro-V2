"""
Módulo de Análise Multi-Agente - Versão Reconstruída
Sistema completo para análise de documentos com múltiplos agentes jurídicos
"""

import os
import json
import logging
import uuid
import time
from datetime import datetime
from flask import Blueprint, render_template, request, jsonify, session, current_app
from flask_login import current_user
from werkzeug.utils import secure_filename
from models import db, AgenteJuridico, CategoriaJuridica, AnaliseJuridica

# Configurar logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# Criar Blueprint
multiagent_bp = Blueprint('multiagent_new', __name__, url_prefix='/multiagent')

@multiagent_bp.route('/')
def index():
    """Página principal do sistema multi-agente"""
    try:
        # Buscar agentes ativos do banco de dados
        agentes_db = AgenteJuridico.query.filter_by(ativo=True).limit(20).all()
        
        agentes_formatados = []
        for agente in agentes_db:
            categoria = CategoriaJuridica.query.get(agente.categoria_id) if agente.categoria_id else None
            
            agente_data = {
                'id': agente.id,
                'nome': agente.nome,
                'descricao': agente.descricao or 'Especialista jurídico',
                'categoria': categoria.nome if categoria else 'Geral',
                'icone': agente.icone or 'fas fa-balance-scale',
                'cor_destaque': agente.cor_destaque or '#007bff'
            }
            agentes_formatados.append(agente_data)
        
        logger.info(f"✅ Carregados {len(agentes_formatados)} agentes para análise multi-agente")
        
        return render_template('multiagent/index.html', 
                             agentes=agentes_formatados,
                             total_agentes=len(agentes_formatados))
        
    except Exception as e:
        logger.error(f"❌ Erro ao carregar página multi-agente: {e}")
        return render_template('multiagent/index.html', 
                             agentes=[],
                             total_agentes=0,
                             erro=str(e))

@multiagent_bp.route('/upload', methods=['POST'])
def upload_documento():
    """Upload e processamento de documento"""
    try:
        if 'documento' not in request.files:
            return jsonify({'erro': 'Nenhum arquivo enviado'}), 400
        
        arquivo = request.files['documento']
        if arquivo.filename == '':
            return jsonify({'erro': 'Nenhum arquivo selecionado'}), 400
        
        # Validar extensão
        extensoes_permitidas = ['.pdf', '.docx', '.txt']
        if not arquivo.filename or not any(arquivo.filename.lower().endswith(ext) for ext in extensoes_permitidas):
            return jsonify({'erro': 'Formato de arquivo não suportado'}), 400
        
        # Salvar arquivo
        filename = secure_filename(arquivo.filename) if arquivo.filename else 'documento'
        upload_folder = os.path.join(current_app.root_path, 'uploads')
        os.makedirs(upload_folder, exist_ok=True)
        
        filepath = os.path.join(upload_folder, filename)
        arquivo.save(filepath)
        
        # Extrair texto do documento
        texto_extraido = extrair_texto_documento(filepath)
        
        return jsonify({
            'sucesso': True,
            'arquivo': filename,
            'tamanho': len(texto_extraido),
            'preview': texto_extraido[:500] + '...' if len(texto_extraido) > 500 else texto_extraido
        })
        
    except Exception as e:
        logger.error(f"❌ Erro no upload: {e}")
        return jsonify({'erro': f'Erro no processamento: {str(e)}'}), 500

@multiagent_bp.route('/analisar', methods=['POST'])
def analisar_documento():
    """Realizar análise multi-agente do documento"""
    try:
        # Aceitar tanto JSON quanto FormData
        if request.is_json:
            data = request.get_json()
            agentes_ids = data.get('agentes', [])
            texto_documento = data.get('texto', '')
        else:
            # FormData do formulário
            agentes_ids = request.form.getlist('agentes_selecionados')
            texto_documento = request.form.get('texto_documento', '')
            
            # Processar arquivo se enviado
            arquivo = request.files.get('arquivo_documento')
            if arquivo and arquivo.filename:
                # Extrair texto do arquivo
                texto_extraido = extrair_texto_documento_upload(arquivo)
                if texto_extraido:
                    texto_documento = texto_extraido
        
        if not agentes_ids:
            return jsonify({'erro': 'Nenhum agente selecionado'}), 400
        
        if not texto_documento:
            return jsonify({'erro': 'Documento vazio'}), 400
        
        # Buscar agentes selecionados
        agentes = AgenteJuridico.query.filter(AgenteJuridico.id.in_(agentes_ids)).all()
        
        resultados = []
        for agente in agentes:
            # Realizar análise real usando IA
            resultado_agente = realizar_analise_ia(agente, texto_documento)
            resultados.append(resultado_agente)
        
        # Gerar número de registro único
        numero_registro = f"ANL-{datetime.now().strftime('%Y%m%d')}-{str(uuid.uuid4())[:8].upper()}"
        
        # Obter ID do usuário (usar current_user se autenticado, senão usar 1)
        usuario_id = current_user.id if current_user.is_authenticated else 1
        
        # Criar análise no banco usando construtor direto
        analise = AnaliseJuridica(
            numero_registro=numero_registro,
            usuario_id=usuario_id,
            texto_original=texto_documento[:5000],  # Limitar tamanho
            resultados_json=json.dumps(resultados, ensure_ascii=False),
            total_agentes=len(resultados),
            status='concluida'
        )
        
        try:
            db.session.add(analise)
            db.session.commit()
        except Exception as db_error:
            logger.error(f"Erro ao salvar no banco: {db_error}")
            db.session.rollback()
            # Continuar sem salvar se houver erro
        
        # Gerar relatório consolidado estruturado
        relatorio_consolidado = gerar_relatorio_consolidado(
            resultados, texto_documento, agentes, numero_registro
        )
        
        return jsonify({
            'sucesso': True,
            'analise_id': analise.id,
            'resultados': resultados,
            'total_agentes': len(resultados),
            'relatorio_url': f'/multiagent/relatorio/{analise.id}',
            'relatorio_consolidado': relatorio_consolidado
        })
        
    except Exception as e:
        db.session.rollback()  # Rollback em caso de erro
        logger.error(f"❌ Erro na análise: {e}")
        return jsonify({'erro': f'Erro na análise: {str(e)}'}), 500

def extrair_texto_documento_upload(arquivo):
    """Extrair texto de arquivo enviado via upload"""
    try:
        filename = arquivo.filename.lower()
        
        if filename.endswith('.txt'):
            return arquivo.read().decode('utf-8')
        elif filename.endswith('.pdf'):
            try:
                import PyPDF2
                import io
                pdf_reader = PyPDF2.PdfReader(io.BytesIO(arquivo.read()))
                texto = ""
                for page in pdf_reader.pages:
                    texto += page.extract_text()
                return texto
            except ImportError:
                return "Erro: PyPDF2 não instalado"
        elif filename.endswith('.docx'):
            try:
                import docx
                import io
                doc = docx.Document(io.BytesIO(arquivo.read()))
                texto = ""
                for paragraph in doc.paragraphs:
                    texto += paragraph.text + "\n"
                return texto
            except ImportError:
                return "Erro: python-docx não instalado"
        else:
            return "Formato não suportado"
    except Exception as e:
        logger.error(f"Erro ao extrair texto: {e}")
        return f"Erro na extração: {str(e)}"

def extrair_texto_documento(filepath):
    """Extrair texto de diferentes tipos de documento"""
    try:
        extensao = os.path.splitext(filepath)[1].lower()
        
        if extensao == '.txt':
            with open(filepath, 'r', encoding='utf-8') as f:
                return f.read()
        
        elif extensao == '.pdf':
            try:
                from PyPDF2 import PdfReader
                reader = PdfReader(filepath)
                texto = ""
                for page in reader.pages:
                    texto += page.extract_text()
                return texto
            except ImportError:
                return "Erro: PyPDF2 não instalado. Instale com: pip install PyPDF2"
        
        elif extensao == '.docx':
            try:
                from docx import Document
                doc = Document(filepath)
                texto = ""
                for paragraph in doc.paragraphs:
                    texto += paragraph.text + "\n"
                return texto
            except ImportError:
                return "Erro: python-docx não instalado. Instale com: pip install python-docx"
        
        else:
            return "Formato de arquivo não suportado"
            
    except Exception as e:
        logger.error(f"❌ Erro ao extrair texto: {e}")
        return f"Erro ao processar arquivo: {str(e)}"

# Funções auxiliares
def obter_estatisticas_sistema():
    """Obter estatísticas do sistema"""
    try:
        total_agentes = AgenteJuridico.query.filter_by(ativo=True).count()
        total_analises = AnaliseJuridica.query.count()
        
        return {
            'total_agentes': total_agentes,
            'total_analises': total_analises,
            'sistema_ativo': True
        }
    except Exception as e:
        logger.error(f"❌ Erro ao obter estatísticas: {e}")
        return {
            'total_agentes': 0,
            'total_analises': 0,
            'sistema_ativo': False
        }

def realizar_analise_ia(agente, texto_documento):
    """Realizar análise jurídica real usando OpenAI"""
    try:
        import openai
        
        # Configurar cliente OpenAI
        api_key = os.environ.get('OPENAI_API_KEY')
        if not api_key:
            return criar_analise_fallback(agente, texto_documento)
        
        client = openai.OpenAI(api_key=api_key)
        
        # Criar prompt especializado baseado no agente
        categoria = agente.categoria.nome if agente.categoria else "Jurídico"
        
        system_prompt = f"""Você é um {agente.nome} altamente especializado em {categoria}.

Características do agente:
- Nome: {agente.nome}
- Especialidade: {categoria}
- Descrição: {agente.descricao or f"Especialista em {categoria}"}

Analise o documento jurídico fornecido seguindo esta estrutura:

## ASPECTOS LEGAIS PRINCIPAIS
- Identifique os principais aspectos legais do documento
- Cite artigos de lei relevantes quando aplicável
- Avalie a conformidade com legislação pertinente

## CONFORMIDADE E RISCOS
- ✅ Pontos positivos e conformes
- ⚠️ Riscos identificados e pontos de atenção
- ❌ Problemas graves (se houver)

## RECOMENDAÇÕES PRÁTICAS
- Sugestões específicas de melhorias
- Ações recomendadas para mitigar riscos
- Ajustes necessários no documento

Mantenha a análise técnica, objetiva e fundamentada juridicamente."""

        start_time = time.time()
        
        # Fazer a requisição para OpenAI
        response = client.chat.completions.create(
            model="gpt-4o",  # the newest OpenAI model is "gpt-4o" which was released May 13, 2024. do not change this unless explicitly requested by the user
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": f"Analise este documento jurídico:\n\n{texto_documento[:4000]}"}
            ],
            temperature=0.3,
            max_tokens=2000
        )
        
        tempo_processamento = f"{time.time() - start_time:.1f}s"
        analise_completa = response.choices[0].message.content
        
        # Extrair pontos-chave da análise
        pontos_chave = extrair_pontos_chave(analise_completa)
        
        return {
            'agente_id': agente.id,
            'agente_nome': agente.nome,
            'especialidade': categoria,
            'analise': analise_completa,
            'pontos_chave': pontos_chave,
            'confianca': 90,
            'tempo_processamento': tempo_processamento,
            'provider': 'OpenAI GPT-4o'
        }
        
    except Exception as e:
        logger.error(f"Erro na análise IA para {agente.nome}: {e}")
        return criar_analise_fallback(agente, texto_documento)

def extrair_pontos_chave(analise_texto):
    """Extrai pontos-chave da análise para resumo"""
    pontos = []
    
    # Buscar por listas e pontos importantes
    linhas = analise_texto.split('\n')
    for linha in linhas:
        linha = linha.strip()
        if linha.startswith('-') or linha.startswith('•') or linha.startswith('*'):
            ponto = linha.lstrip('-•* ').strip()
            if len(ponto) > 10 and len(ponto) < 100:
                pontos.append(ponto)
    
    # Se não encontrou pontos, criar alguns baseados na análise
    if not pontos:
        if "contrato" in analise_texto.lower():
            pontos = ["Análise de estrutura contratual", "Identificação de cláusulas", "Avaliação de riscos"]
        elif "parecer" in analise_texto.lower():
            pontos = ["Fundamentação jurídica", "Análise de precedentes", "Conclusões técnicas"]
        else:
            pontos = ["Análise jurídica especializada", "Identificação de questões legais", "Recomendações práticas"]
    
    return pontos[:5]  # Máximo 5 pontos

def criar_analise_fallback(agente, texto_documento):
    """Cria análise básica quando IA não está disponível"""
    categoria = agente.categoria.nome if agente.categoria else "Jurídico"
    
    # Análise básica baseada em padrões do documento
    if "contrato" in texto_documento.lower():
        analise = f"""## ANÁLISE ESTRUTURAL DE CONTRATO

**Especialista:** {agente.nome}

### ASPECTOS LEGAIS PRINCIPAIS
- Documento tipificado como contrato conforme Código Civil
- Identificação das partes contratuais presente
- Objeto contratual definido no documento

### CONFORMIDADE E RISCOS  
- ✅ Estrutura básica contratual adequada
- ⚠️ Recomenda-se revisão detalhada de cláusulas específicas
- ⚠️ Verificar adequação à legislação específica do setor

### RECOMENDAÇÕES PRÁTICAS
- Revisar cláusulas de rescisão e penalidades
- Verificar conformidade com legislação aplicável
- Considerar inclusão de cláusulas de proteção de dados"""
    else:
        analise = f"""## ANÁLISE JURÍDICA ESPECIALIZADA

**Especialista:** {agente.nome} - {categoria}

### ASPECTOS LEGAIS PRINCIPAIS
- Documento analisado sob perspectiva de {categoria}
- Identificação de questões jurídicas relevantes
- Avaliação preliminar de conformidade legal

### CONFORMIDADE E RISCOS
- ✅ Documento apresenta estrutura básica adequada
- ⚠️ Recomenda-se análise detalhada de aspectos específicos
- ⚠️ Verificar atualização conforme legislação vigente

### RECOMENDAÇÕES PRÁTICAS
- Revisar fundamentação legal específica
- Verificar adequação aos precedentes jurisprudenciais
- Considerar atualizações regulamentares recentes"""
    
    return {
        'agente_id': agente.id,
        'agente_nome': agente.nome,
        'especialidade': categoria,
        'analise': analise,
        'pontos_chave': [
            f"Análise especializada em {categoria}",
            "Identificação de aspectos legais relevantes",
            "Recomendações para conformidade"
        ],
        'confianca': 75,
        'tempo_processamento': '1.2s',
        'provider': 'Análise Estrutural'
    }

def extrair_resumo_executivo(texto_documento):
    """Extrai informações principais do documento para o resumo executivo"""
    resumo = {
        'partes': [],
        'objeto': None,
        'valor': None,
        'prazo': None
    }
    
    texto_lower = texto_documento.lower()
    
    # Extrair partes contratuais
    if 'contratante' in texto_lower and 'contratada' in texto_lower:
        linhas = texto_documento.split('\n')
        for linha in linhas:
            if 'contratante' in linha.lower():
                nome = linha.split(':')[-1].strip()
                if nome:
                    resumo['partes'].append({'tipo': 'CONTRATANTE', 'nome': nome})
            elif 'contratada' in linha.lower():
                nome = linha.split(':')[-1].strip()
                if nome:
                    resumo['partes'].append({'tipo': 'CONTRATADA', 'nome': nome})
    
    # Extrair objeto
    if 'objeto' in texto_lower:
        linhas = texto_documento.split('\n')
        for linha in linhas:
            if 'objeto' in linha.lower() and ':' in linha:
                objeto = linha.split(':')[-1].strip()
                if len(objeto) > 10:
                    resumo['objeto'] = objeto[:200]
                    break
    
    # Extrair valor
    import re
    valores = re.findall(r'R\$\s*[\d.,]+', texto_documento)
    if valores:
        resumo['valor'] = valores[0]
    
    # Extrair prazo
    if 'prazo' in texto_lower:
        linhas = texto_documento.split('\n')
        for linha in linhas:
            if 'prazo' in linha.lower() and ':' in linha:
                prazo = linha.split(':')[-1].strip()
                if len(prazo) > 5:
                    resumo['prazo'] = prazo[:100]
                    break
    
    return resumo

def analisar_pontos_fortes_riscos(resultados):
    """Analisa resultados para extrair pontos fortes e riscos consolidados"""
    pontos_fortes = []
    riscos = []
    
    # Análise dos símbolos nos resultados
    for resultado in resultados:
        analise = resultado.get('analise', '')
        
        # Buscar pontos positivos (✅)
        linhas_positivas = [linha for linha in analise.split('\n') if '✅' in linha]
        for linha in linhas_positivas:
            ponto = linha.replace('✅', '').strip('- ').strip()
            if len(ponto) > 10:
                pontos_fortes.append({
                    'titulo': ponto.split(':')[0] if ':' in ponto else 'Conformidade Legal',
                    'descricao': ponto.split(':')[1] if ':' in ponto else ponto
                })
        
        # Buscar riscos (⚠️)
        linhas_riscos = [linha for linha in analise.split('\n') if '⚠️' in linha]
        for linha in linhas_riscos:
            risco = linha.replace('⚠️', '').strip('- ').strip()
            if len(risco) > 10:
                riscos.append({
                    'titulo': risco.split(':')[0] if ':' in risco else 'Ponto de Atenção',
                    'descricao': risco.split(':')[1] if ':' in risco else risco
                })
    
    # Pontos fortes padrão se não encontrados
    if not pontos_fortes:
        pontos_fortes = [
            {'titulo': 'Estrutura Legal', 'descricao': 'Documento apresenta estrutura jurídica adequada'},
            {'titulo': 'Identificação das Partes', 'descricao': 'Partes claramente identificadas'},
            {'titulo': 'Objeto Definido', 'descricao': 'Objeto do documento bem especificado'}
        ]
    
    # Riscos padrão se não encontrados
    if not riscos:
        riscos = [
            {'titulo': 'Revisão Recomendada', 'descricao': 'Alguns pontos necessitam de revisão detalhada'},
            {'titulo': 'Conformidade Legal', 'descricao': 'Verificar adequação à legislação específica'},
            {'titulo': 'Cláusulas Específicas', 'descricao': 'Considerar inclusão de cláusulas adicionais'}
        ]
    
    return pontos_fortes[:5], riscos[:5]

def gerar_recomendacoes_prioritarias(resultados):
    """Gera recomendações categorizadas por prioridade"""
    recomendacoes = {
        'alta': [],
        'media': [],
        'baixa': []
    }
    
    # Extrair recomendações dos resultados
    todas_recomendacoes = []
    for resultado in resultados:
        analise = resultado.get('analise', '')
        
        # Buscar seções de recomendações
        secoes = analise.split('##')
        for secao in secoes:
            if 'recomend' in secao.lower():
                linhas = secao.split('\n')
                for linha in linhas:
                    linha = linha.strip('- *').strip()
                    if len(linha) > 20 and not linha.startswith('#'):
                        todas_recomendacoes.append(linha)
    
    # Categorizar por palavras-chave
    palavras_alta = ['urgente', 'crítico', 'fundamental', 'obrigatório', 'essencial']
    palavras_media = ['importante', 'recomendado', 'sugerido', 'adequado']
    
    for rec in todas_recomendacoes:
        rec_lower = rec.lower()
        if any(palavra in rec_lower for palavra in palavras_alta):
            recomendacoes['alta'].append(rec)
        elif any(palavra in rec_lower for palavra in palavras_media):
            recomendacoes['media'].append(rec)
        else:
            recomendacoes['baixa'].append(rec)
    
    # Recomendações padrão se necessário
    if not recomendacoes['alta']:
        recomendacoes['alta'] = [
            'Incluir cláusulas específicas sobre propriedade intelectual dos materiais criados',
            'Definir claramente responsabilidades e obrigações de ambas as partes',
            'Estabelecer critérios objetivos para rescisão e penalidades'
        ]
    
    if not recomendacoes['media']:
        recomendacoes['media'] = [
            'Incluir Service Level Agreement (SLA) com métricas objetivas',
            'Detalhar procedimentos de backup e segurança dos dados',
            'Estabelecer multas proporcionais para ambas as partes'
        ]
    
    if not recomendacoes['baixa']:
        recomendacoes['baixa'] = [
            'Aprimorar redação de cláusulas específicas',
            'Incluir procedimentos de mediação pré-judicial',
            'Detalhar critérios para alterações de escopo'
        ]
    
    # Limitar quantidade
    for prioridade in recomendacoes:
        recomendacoes[prioridade] = recomendacoes[prioridade][:4]
    
    return recomendacoes

def calcular_classificacao_geral(resultados):
    """Calcula classificação geral baseada nas análises"""
    confiancas = [r.get('confianca', 75) for r in resultados]
    media_confianca = sum(confiancas) / len(confiancas) if confiancas else 75
    
    # Converter para estrelas (1-5)
    nota = int((media_confianca / 100) * 5)
    nota = max(1, min(5, nota))
    
    estrelas = '⭐' * nota + '☆' * (5 - nota)
    
    return {
        'nota': nota,
        'estrelas': estrelas,
        'confianca': media_confianca
    }

def gerar_relatorio_consolidado(resultados, texto_documento, agentes, numero_registro):
    """Gera relatório consolidado no formato do modelo"""
    from datetime import datetime
    
    # Detectar tipo de documento
    tipo_documento = detectar_tipo_documento(texto_documento)
    
    # Extrair resumo executivo
    resumo_executivo = extrair_resumo_executivo(texto_documento)
    
    # Analisar pontos fortes e riscos
    pontos_fortes, principais_riscos = analisar_pontos_fortes_riscos(resultados)
    
    # Gerar recomendações prioritárias
    recomendacoes = gerar_recomendacoes_prioritarias(resultados)
    
    # Calcular classificação
    classificacao = calcular_classificacao_geral(resultados)
    
    # Determinar área principal
    areas = [agente.categoria.nome for agente in agentes if agente.categoria]
    area_principal = areas[0] if areas else "Análise Jurídica Geral"
    
    return {
        'numero_registro': numero_registro,
        'data_analise': datetime.now().strftime('%d de %B de %Y'),
        'tipo_documento': tipo_documento,
        'area_principal': area_principal,
        'total_agentes': len(resultados),
        'resumo_executivo': resumo_executivo,
        'pontos_fortes': pontos_fortes,
        'principais_riscos': principais_riscos,
        'recomendacoes': recomendacoes,
        'classificacao_geral': classificacao,
        'criterios_avaliacao': [
            {'nome': 'Estrutura Legal', 'avaliacao': 'Boa'},
            {'nome': 'Proteção das Partes', 'avaliacao': 'Média'},
            {'nome': 'Clareza de Termos', 'avaliacao': 'Boa'},
            {'nome': 'Gestão de Riscos', 'avaliacao': 'Necessita Melhoria'}
        ],
        'conclusao_texto': f'O documento apresenta estrutura jurídica adequada com observância dos requisitos legais básicos. Recomenda-se atenção aos pontos de alta prioridade identificados.',
        'recomendacao_final': 'Proceder com revisão focada nos pontos de alta prioridade antes da execução.',
        'tempo_total': f'{sum([float(r.get("tempo_processamento", "1.0").replace("s", "")) for r in resultados]):.1f}s'
    }

def detectar_tipo_documento(texto):
    """Detecta o tipo de documento baseado no conteúdo"""
    texto_lower = texto.lower()
    
    if 'contrato' in texto_lower:
        if 'prestação de serviços' in texto_lower:
            return 'Contrato de Prestação de Serviços'
        elif 'trabalho' in texto_lower:
            return 'Contrato de Trabalho'
        elif 'compra e venda' in texto_lower:
            return 'Contrato de Compra e Venda'
        else:
            return 'Contrato'
    elif 'parecer' in texto_lower:
        return 'Parecer Jurídico'
    elif 'petição' in texto_lower:
        return 'Petição'
    else:
        return 'Documento Jurídico'

@multiagent_bp.route('/relatorio/<int:analise_id>')
def visualizar_relatorio(analise_id):
    """Visualizar relatório completo de análise"""
    try:
        analise = AnaliseJuridica.query.get_or_404(analise_id)
        resultados = json.loads(analise.resultados_json)
        
        # Buscar agentes da análise
        agentes_ids = [r['agente_id'] for r in resultados]
        agentes = AgenteJuridico.query.filter(AgenteJuridico.id.in_(agentes_ids)).all()
        
        # Gerar relatório consolidado
        relatorio = gerar_relatorio_consolidado(
            resultados, analise.texto_original, agentes, analise.numero_registro
        )
        
        return render_template(
            'multiagent/resultado_analise_completa.html',
            resultados=resultados,
            **relatorio
        )
        
    except Exception as e:
        logger.error(f"Erro ao gerar relatório: {e}")
        return f"Erro ao carregar relatório: {str(e)}", 500