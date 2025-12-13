"""
API de exportação para transcrições de vídeo
Suporta múltiplos formatos preservando a estrutura dos dados
"""
import json
import csv
import io
from datetime import datetime
from flask import request, jsonify, send_file
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.shared import OxmlElement, qn
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_JUSTIFY


def register_video_export_api(app):
    """Registra as rotas de exportação de vídeo"""
    
    @app.route('/video/export', methods=['POST'])
    def export_transcription():
        """Exporta transcrição em diferentes formatos preservando estrutura"""
        try:
            data = request.get_json()
            if not data:
                return jsonify({'error': 'Dados não fornecidos'}), 400
            
            # Verificar se é exportação direta via transcript_id
            if data.get('export_type') == 'direct' and data.get('transcript_id'):
                return export_direct_from_transcript(data.get('transcript_id'), data.get('format', 'txt'))
            
            format_type = data.get('metadata', {}).get('format', data.get('format', 'txt'))
            content = data.get('content', {})
            metadata = data.get('metadata', {})
            
            # Gerar o arquivo no formato solicitado
            if format_type == 'txt':
                return export_txt(content, metadata)
            elif format_type == 'docx':
                return export_docx(content, metadata)
            elif format_type == 'pdf':
                return export_pdf(content, metadata)
            elif format_type == 'json':
                return export_json(content, metadata)
            elif format_type == 'csv':
                return export_csv(content, metadata)
            elif format_type == 'srt':
                return export_srt(content, metadata)
            else:
                return jsonify({'error': f'Formato {format_type} não suportado'}), 400
                
        except Exception as e:
            print(f"Erro na exportação: {str(e)}")
            return jsonify({'error': f'Erro na exportação: {str(e)}'}), 500

    def export_txt(content, metadata):
        """Exporta em formato TXT simples"""
        output = io.StringIO()
        
        # Cabeçalho
        output.write("=" * 60 + "\n")
        output.write("TRANSCRIÇÃO DE VÍDEO\n")
        output.write("=" * 60 + "\n\n")
        
        output.write(f"Arquivo: {metadata.get('filename', 'N/A')}\n")
        output.write(f"Data da Exportação: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}\n")
        output.write(f"Duração: {format_duration(metadata.get('duration', 0))}\n")
        output.write(f"Confiança Média: {metadata.get('confidence', 0) * 100:.1f}%\n")
        # Calcular falantes únicos se não fornecido
        unique_speakers = metadata.get('speaker_count', 0)
        if unique_speakers == 0 and content.get('speaker_segments'):
            unique_speakers = len(set(seg.get('speaker', '') for seg in content['speaker_segments']))
        
        output.write(f"Número de Falantes: {unique_speakers}\n\n")
        
        output.write("-" * 60 + "\n")
        output.write("TRANSCRIÇÃO COMPLETA\n")
        output.write("-" * 60 + "\n\n")
        
        # Conteúdo da transcrição
        if content.get('speaker_segments'):
            for segment in content['speaker_segments']:
                timestamp = format_timestamp(segment.get('start', 0))
                speaker = segment.get('speaker', 'Falante Desconhecido')
                text = segment.get('text', '')
                confidence = segment.get('confidence', 0) * 100
                
                output.write(f"[{timestamp}] {speaker} ({confidence:.1f}%)\n")
                output.write(f"{text}\n\n")
        else:
            output.write(content.get('text', 'Transcrição não disponível'))
        
        # Converter para bytes
        txt_content = output.getvalue().encode('utf-8')
        output.close()
        
        return send_file(
            io.BytesIO(txt_content),
            mimetype='text/plain',
            as_attachment=True,
            download_name=f"{metadata.get('filename', 'transcricao')}.txt"
        )

    def export_docx(content, metadata):
        """Exporta em formato DOCX com formatação profissional"""
        doc = Document()
        
        # Configurar estilos
        styles = doc.styles
        
        # Título principal
        title = doc.add_heading('TRANSCRIÇÃO DE VÍDEO', level=1)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Informações do arquivo
        info_table = doc.add_table(rows=5, cols=2)
        info_table.style = 'Table Grid'
        
        info_data = [
            ['Arquivo:', metadata.get('filename', 'N/A')],
            ['Data da Exportação:', datetime.now().strftime('%d/%m/%Y %H:%M:%S')],
            ['Duração:', format_duration(metadata.get('duration', 0))],
            ['Confiança Média:', f"{metadata.get('confidence', 0) * 100:.1f}%"],
            ['Número de Falantes:', str(metadata.get('speaker_count', 0))]
        ]
        
        for i, (label, value) in enumerate(info_data):
            info_table.rows[i].cells[0].text = label
            info_table.rows[i].cells[1].text = value
            info_table.rows[i].cells[0].paragraphs[0].runs[0].bold = True
        
        # Título da transcrição
        doc.add_heading('TRANSCRIÇÃO DETALHADA', level=2)
        
        # Conteúdo da transcrição
        if content.get('speaker_segments'):
            for i, segment in enumerate(content['speaker_segments']):
                timestamp = format_timestamp(segment.get('start', 0))
                speaker = segment.get('speaker', 'Falante Desconhecido')
                text = segment.get('text', '')
                confidence = segment.get('confidence', 0) * 100
                
                # Cabeçalho do segmento
                header_p = doc.add_paragraph()
                header_p.add_run(f"[{timestamp}] ").bold = True
                header_p.add_run(f"{speaker} ").bold = True
                header_p.add_run(f"(Confiança: {confidence:.1f}%)")
                
                # Texto do segmento
                text_p = doc.add_paragraph(text)
                text_p.style = 'Normal'
                
                # Espaçamento entre segmentos
                if i < len(content['speaker_segments']) - 1:
                    doc.add_paragraph()
        else:
            doc.add_paragraph(content.get('text', 'Transcrição não disponível'))
        
        # Salvar em bytes
        docx_buffer = io.BytesIO()
        doc.save(docx_buffer)
        docx_buffer.seek(0)
        
        return send_file(
            docx_buffer,
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            as_attachment=True,
            download_name=f"{metadata.get('filename', 'transcricao')}.docx"
        )

    def export_pdf(content, metadata):
        """Exporta em formato PDF profissional"""
        pdf_buffer = io.BytesIO()
        doc = SimpleDocTemplate(pdf_buffer, pagesize=A4, rightMargin=72, leftMargin=72, topMargin=72, bottomMargin=18)
        
        # Estilos
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=20,
            spaceAfter=30,
            alignment=TA_CENTER,
            textColor=colors.darkblue
        )
        
        header_style = ParagraphStyle(
            'CustomHeader',
            parent=styles['Heading2'],
            fontSize=14,
            spaceAfter=12,
            textColor=colors.darkgreen
        )
        
        speaker_style = ParagraphStyle(
            'SpeakerStyle',
            parent=styles['Normal'],
            fontSize=10,
            spaceBefore=6,
            spaceAfter=3,
            textColor=colors.blue,
            fontName='Helvetica-Bold'
        )
        
        content_style = ParagraphStyle(
            'ContentStyle',
            parent=styles['Normal'],
            fontSize=11,
            spaceAfter=12,
            alignment=TA_JUSTIFY,
            leftIndent=20
        )
        
        # Construir conteúdo
        story = []
        
        # Título
        story.append(Paragraph("TRANSCRIÇÃO DE VÍDEO", title_style))
        story.append(Spacer(1, 20))
        
        # Tabela de informações
        info_data = [
            ['Arquivo:', metadata.get('filename', 'N/A')],
            ['Data da Exportação:', datetime.now().strftime('%d/%m/%Y %H:%M:%S')],
            ['Duração:', format_duration(metadata.get('duration', 0))],
            ['Confiança Média:', f"{metadata.get('confidence', 0) * 100:.1f}%"],
            ['Número de Falantes:', str(metadata.get('speaker_count', 0))]
        ]
        
        info_table = Table(info_data, colWidths=[2*inch, 3*inch])
        info_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.lightgrey),
            ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ]))
        
        story.append(info_table)
        story.append(Spacer(1, 30))
        
        # Título da transcrição
        story.append(Paragraph("TRANSCRIÇÃO DETALHADA", header_style))
        story.append(Spacer(1, 15))
        
        # Conteúdo da transcrição
        if content.get('speaker_segments'):
            for segment in content['speaker_segments']:
                timestamp = format_timestamp(segment.get('start', 0))
                speaker = segment.get('speaker', 'Falante Desconhecido')
                text = segment.get('text', '')
                confidence = segment.get('confidence', 0) * 100
                
                # Cabeçalho do segmento
                speaker_header = f"[{timestamp}] {speaker} (Confiança: {confidence:.1f}%)"
                story.append(Paragraph(speaker_header, speaker_style))
                
                # Texto do segmento
                story.append(Paragraph(text, content_style))
                story.append(Spacer(1, 8))
        else:
            story.append(Paragraph(content.get('text', 'Transcrição não disponível'), content_style))
        
        # Gerar PDF
        doc.build(story)
        pdf_buffer.seek(0)
        
        return send_file(
            pdf_buffer,
            mimetype='application/pdf',
            as_attachment=True,
            download_name=f"{metadata.get('filename', 'transcricao')}.pdf"
        )

    def export_json(content, metadata):
        """Exporta em formato JSON estruturado"""
        export_data = {
            'metadata': {
                'filename': metadata.get('filename', 'N/A'),
                'export_date': datetime.now().isoformat(),
                'duration_ms': metadata.get('duration', 0),
                'duration_formatted': format_duration(metadata.get('duration', 0)),
                'confidence': metadata.get('confidence', 0),
                'speaker_count': metadata.get('speaker_count', 0),
                'transcript_id': metadata.get('transcript_id'),
                'format_version': '1.0'
            },
            'transcription': {
                'full_text': content.get('text', ''),
                'segments': []
            },
            'analysis': {
                'word_count': len(content.get('text', '').split()) if content.get('text') else 0,
                'confidence_distribution': calculate_confidence_distribution(content),
                'speaker_distribution': calculate_speaker_distribution(content)
            }
        }
        
        # Processar segmentos
        if content.get('speaker_segments'):
            for i, segment in enumerate(content['speaker_segments']):
                segment_data = {
                    'id': i + 1,
                    'speaker': segment.get('speaker', 'Falante Desconhecido'),
                    'start_ms': segment.get('start', 0),
                    'end_ms': segment.get('end', 0),
                    'start_formatted': format_timestamp(segment.get('start', 0)),
                    'end_formatted': format_timestamp(segment.get('end', 0)),
                    'duration_ms': segment.get('end', 0) - segment.get('start', 0),
                    'text': segment.get('text', ''),
                    'confidence': segment.get('confidence', 0),
                    'word_count': len(segment.get('text', '').split())
                }
                export_data['transcription']['segments'].append(segment_data)
        
        json_content = json.dumps(export_data, indent=2, ensure_ascii=False)
        
        return send_file(
            io.BytesIO(json_content.encode('utf-8')),
            mimetype='application/json',
            as_attachment=True,
            download_name=f"{metadata.get('filename', 'transcricao')}.json"
        )

    def export_csv(content, metadata):
        """Exporta em formato CSV para análise em planilhas"""
        csv_buffer = io.StringIO()
        writer = csv.writer(csv_buffer)
        
        # Cabeçalho
        writer.writerow([
            'ID',
            'Falante',
            'Início (ms)',
            'Fim (ms)',
            'Início (formatado)',
            'Fim (formatado)',
            'Duração (ms)',
            'Texto',
            'Confiança (%)',
            'Palavras'
        ])
        
        # Dados dos segmentos
        if content.get('speaker_segments'):
            for i, segment in enumerate(content['speaker_segments']):
                writer.writerow([
                    i + 1,
                    segment.get('speaker', 'Falante Desconhecido'),
                    segment.get('start', 0),
                    segment.get('end', 0),
                    format_timestamp(segment.get('start', 0)),
                    format_timestamp(segment.get('end', 0)),
                    segment.get('end', 0) - segment.get('start', 0),
                    segment.get('text', ''),
                    f"{segment.get('confidence', 0) * 100:.1f}",
                    len(segment.get('text', '').split())
                ])
        
        csv_content = csv_buffer.getvalue().encode('utf-8')
        csv_buffer.close()
        
        return send_file(
            io.BytesIO(csv_content),
            mimetype='text/csv',
            as_attachment=True,
            download_name=f"{metadata.get('filename', 'transcricao')}.csv"
        )

    def export_srt(content, metadata):
        """Exporta em formato SRT para legendas"""
        srt_buffer = io.StringIO()
        
        if content.get('speaker_segments'):
            for i, segment in enumerate(content['speaker_segments']):
                start_time = format_srt_timestamp(segment.get('start', 0))
                end_time = format_srt_timestamp(segment.get('end', 0))
                speaker = segment.get('speaker', 'Falante')
                text = segment.get('text', '')
                
                srt_buffer.write(f"{i + 1}\n")
                srt_buffer.write(f"{start_time} --> {end_time}\n")
                srt_buffer.write(f"[{speaker}] {text}\n\n")
        
        srt_content = srt_buffer.getvalue().encode('utf-8')
        srt_buffer.close()
        
        return send_file(
            io.BytesIO(srt_content),
            mimetype='text/plain',
            as_attachment=True,
            download_name=f"{metadata.get('filename', 'transcricao')}.srt"
        )

    def format_duration(duration_ms):
        """Formata duração em milissegundos para string legível"""
        if not duration_ms:
            return "0:00"
        
        seconds = int(duration_ms / 1000)
        minutes = seconds // 60
        seconds = seconds % 60
        hours = minutes // 60
        minutes = minutes % 60
        
        if hours > 0:
            return f"{hours}:{minutes:02d}:{seconds:02d}"
        else:
            return f"{minutes}:{seconds:02d}"

    def format_timestamp(timestamp_ms):
        """Formata timestamp para exibição"""
        if not timestamp_ms:
            return "00:00"
        
        total_seconds = int(timestamp_ms / 1000)
        minutes = total_seconds // 60
        seconds = total_seconds % 60
        
        return f"{minutes:02d}:{seconds:02d}"

    def format_srt_timestamp(timestamp_ms):
        """Formata timestamp para formato SRT"""
        if not timestamp_ms:
            return "00:00:00,000"
        
        total_ms = int(timestamp_ms)
        hours = total_ms // 3600000
        total_ms %= 3600000
        minutes = total_ms // 60000
        total_ms %= 60000
        seconds = total_ms // 1000
        milliseconds = total_ms % 1000
        
        return f"{hours:02d}:{minutes:02d}:{seconds:02d},{milliseconds:03d}"

    def calculate_confidence_distribution(content):
        """Calcula distribuição de confiança"""
        if not content.get('speaker_segments'):
            return {'high': 0, 'medium': 0, 'low': 0}
        
        high = medium = low = 0
        for segment in content['speaker_segments']:
            confidence = segment.get('confidence', 0)
            if confidence >= 0.8:
                high += 1
            elif confidence >= 0.6:
                medium += 1
            else:
                low += 1
        
        return {'high': high, 'medium': medium, 'low': low}

    def calculate_speaker_distribution(content):
        """Calcula distribuição por falante"""
        if not content.get('speaker_segments'):
            return {}
        
        distribution = {}
        for segment in content['speaker_segments']:
            speaker = segment.get('speaker', 'Desconhecido')
            distribution[speaker] = distribution.get(speaker, 0) + 1
        
        return distribution

    def export_direct_from_transcript(transcript_id, format_type):
        """Exporta diretamente usando transcript_id do AssemblyAI"""
        try:
            import assemblyai as aai
            import os
            
            # Configurar AssemblyAI
            aai.settings.api_key = os.getenv('ASSEMBLYAI_API_KEY')
            transcript = aai.Transcript.get_by_id(transcript_id)
            
            if transcript.status != aai.TranscriptStatus.completed:
                return jsonify({'error': 'Transcrição não concluída'}), 400
            
            # Processar dados para formato esperado
            content = {
                'text': transcript.text,
                'speaker_segments': []
            }
            
            if hasattr(transcript, 'utterances') and transcript.utterances:
                for utterance in transcript.utterances:
                    content['speaker_segments'].append({
                        'speaker': utterance.speaker,
                        'text': utterance.text,
                        'start': utterance.start,
                        'end': utterance.end,
                        'confidence': utterance.confidence if hasattr(utterance, 'confidence') else 0.8
                    })
            
            metadata = {
                'filename': f'transcricao_{transcript_id}',
                'duration': transcript.audio_duration,
                'confidence': transcript.confidence if hasattr(transcript, 'confidence') else 0.8,
                'speaker_count': len(set(seg['speaker'] for seg in content['speaker_segments'])) if content['speaker_segments'] else 0
            }
            
            # Gerar arquivo no formato solicitado
            if format_type == 'txt':
                return export_txt(content, metadata)
            elif format_type == 'docx':
                return export_docx(content, metadata)
            elif format_type == 'json':
                return export_json(content, metadata)
            else:
                return jsonify({'error': f'Formato {format_type} não suportado'}), 400
                
        except Exception as e:
            print(f"Erro na exportação direta: {str(e)}")
            return jsonify({'error': f'Erro na exportação: {str(e)}'}), 500

    print("✅ API de exportação de vídeo registrada com sucesso")