"""
Rotas para o módulo de Análise Legal BERTimbau
"""

from flask import render_template, request, jsonify, flash, redirect, url_for, current_app, send_file
from flask_login import login_required, current_user
from werkzeug.utils import secure_filename
from functools import lru_cache
import os
import json
import time
from datetime import datetime

from . import analise_legal_bert_bp
from .models import AnaliseLegalDocumentoBert, EntidadeExtraidaBert, ClassificacaoDocumentoBert
from .services.legal_analyzer import LegalAnalyzer
from main import db

# Inicializar analisador legal com cache
analyzer = None

@lru_cache(maxsize=1)
def get_analyzer():
    """Lazy loading do analisador legal com cache"""
    global analyzer
    if analyzer is None:
        analyzer = LegalAnalyzer()
        current_app.logger.info("✅ LegalAnalyzer inicializado via lazy loading")
    return analyzer


@analise_legal_bert_bp.route('/', strict_slashes=False)
@login_required
def index():
    """Página principal do módulo de análise legal"""
    
    # Buscar documentos recentes do usuário
    documentos_recentes = AnaliseLegalDocumentoBert.query.filter_by(
        usuario_id=current_user.id
    ).order_by(AnaliseLegalDocumentoBert.created_at.desc()).limit(5).all()
    
    # Estatísticas do usuário
    total_documentos = AnaliseLegalDocumentoBert.query.filter_by(usuario_id=current_user.id).count()
    documentos_concluidos = AnaliseLegalDocumentoBert.query.filter_by(
        usuario_id=current_user.id, 
        status_analise='concluida'
    ).count()
    
    # Informações do modelo
    try:
        model_info = get_analyzer().get_model_info()
    except Exception as e:
        current_app.logger.error(f"Erro ao obter info do modelo: {e}")
        model_info = {'nome': 'Legal Analyzer', 'status': 'erro'}
    
    return render_template('analise_legal_bert/index.html',
                         documentos_recentes=documentos_recentes,
                         total_documentos=total_documentos,
                         documentos_concluidos=documentos_concluidos,
                         model_info=model_info)


@analise_legal_bert_bp.route('/nova-analise', methods=['GET', 'POST'])
@login_required
def nova_analise():
    """Página para nova análise de documento"""
    
    if request.method == 'POST':
        try:
            # Obter dados do formulário
            titulo = request.form.get('titulo', '').strip()
            conteudo = request.form.get('conteudo', '').strip()
            
            # Obter configurações de análise
            config_analise = {
                'tipoDocumento': request.form.get('tipo-documento', 'auto'),
                'extrairEntidades': request.form.get('extrairEntidades') == 'on',
                'classificarDocumento': request.form.get('classificarDocumento') == 'on',
                'gerarResumo': request.form.get('gerarResumo') == 'on',
                'detectarProblemas': request.form.get('detectarProblemas') == 'on',
                'analiseSentimento': request.form.get('analiseSentimento') == 'on',
                'buscarSimilaridade': request.form.get('buscarSimilaridade') == 'on'
            }
            
            # Configurações padrão se não especificadas
            if all(not v for v in config_analise.values() if isinstance(v, bool)):
                config_analise = {
                    'tipoDocumento': config_analise['tipoDocumento'],
                    'extrairEntidades': True,
                    'classificarDocumento': True,
                    'gerarResumo': True,
                    'detectarProblemas': True,
                    'analiseSentimento': False,
                    'buscarSimilaridade': False
                }
            
            # Validações
            if not titulo:
                flash('Título é obrigatório', 'error')
                return render_template('analise_legal_bert/nova_analise.html')
            
            if not conteudo or len(conteudo) < 50:
                flash('Conteúdo deve ter pelo menos 50 caracteres', 'error')
                return render_template('analise_legal_bert/nova_analise.html')
            
            # Criar registro no banco
            documento = AnaliseLegalDocumentoBert()
            documento.titulo = titulo
            documento.conteudo_original = conteudo
            documento.usuario_id = current_user.id
            documento.status_analise = 'processando'
            
            db.session.add(documento)
            db.session.commit()
            
            # Processar análise com configurações
            resultado = processar_analise_documento(documento.id, titulo, conteudo, config_analise)
            
            if resultado['sucesso']:
                flash('Análise concluída com sucesso!', 'success')
                return redirect(url_for('analise_legal_bert.resultado', doc_id=documento.id))
            else:
                flash(f'Erro na análise: {resultado["erro"]}', 'error')
                return render_template('analise_legal_bert/nova_analise.html')
                
        except Exception as e:
            current_app.logger.error(f"Erro ao processar nova análise: {e}")
            flash('Erro interno do servidor', 'error')
            return render_template('analise_legal_bert/nova_analise.html')
    
    return render_template('analise_legal_bert/nova_analise.html')


@analise_legal_bert_bp.route('/upload', methods=['POST'])
@login_required
def upload_documento():
    """Upload e análise de arquivo de documento"""
    
    try:
        if 'arquivo' not in request.files:
            return jsonify({'sucesso': False, 'erro': 'Nenhum arquivo enviado'})
        
        arquivo = request.files['arquivo']
        
        if arquivo.filename == '':
            return jsonify({'sucesso': False, 'erro': 'Nenhum arquivo selecionado'})
        
        # Validar extensão
        extensoes_permitidas = {'txt', 'pdf', 'docx', 'doc'}
        filename_safe = arquivo.filename or ''
        extensao = filename_safe.rsplit('.', 1)[1].lower() if '.' in filename_safe else ''
        
        if extensao not in extensoes_permitidas:
            return jsonify({
                'sucesso': False, 
                'erro': f'Extensão não permitida. Use: {", ".join(extensoes_permitidas)}'
            })
        
        # Salvar arquivo temporariamente
        filename = secure_filename(arquivo.filename or 'documento')
        timestamp = int(time.time())
        temp_filename = f"{timestamp}_{filename}"
        
        upload_dir = os.path.join(current_app.root_path, 'temp_uploads')
        os.makedirs(upload_dir, exist_ok=True)
        
        filepath = os.path.join(upload_dir, temp_filename)
        arquivo.save(filepath)
        
        # Extrair texto do arquivo
        try:
            if extensao == 'txt':
                with open(filepath, 'r', encoding='utf-8') as f:
                    conteudo = f.read()
            elif extensao == 'pdf':
                conteudo = extrair_texto_pdf(filepath)
            elif extensao in ['docx', 'doc']:
                conteudo = extrair_texto_docx(filepath)
            else:
                raise ValueError(f"Extensão {extensao} não suportada")
            
            # Limpar arquivo temporário
            os.remove(filepath)
            
            if not conteudo or len(conteudo.strip()) < 20:
                return jsonify({
                    'sucesso': False, 
                    'erro': 'Arquivo vazio ou com conteúdo insuficiente'
                })
            
            # Criar documento no banco
            documento = AnaliseLegalDocumentoBert()
            documento.titulo = filename
            documento.conteudo_original = conteudo
            documento.usuario_id = current_user.id
            documento.status_analise = 'processando'
            
            db.session.add(documento)
            db.session.commit()
            
            # Configurações padrão para upload
            config_analise = {
                'tipoDocumento': 'auto',
                'extrairEntidades': True,
                'classificarDocumento': True,
                'gerarResumo': True,
                'detectarProblemas': True,
                'analiseSentimento': False,
                'buscarSimilaridade': False
            }
            
            # Processar análise com configurações padrão
            resultado = processar_analise_documento(documento.id, filename, conteudo, config_analise)
            
            if resultado['sucesso']:
                return jsonify({
                    'sucesso': True,
                    'documento_id': documento.id,
                    'redirect_url': url_for('analise_legal_bert.resultado', doc_id=documento.id)
                })
            else:
                return jsonify({'sucesso': False, 'erro': resultado['erro']})
                
        except Exception as e:
            # Limpar arquivo em caso de erro
            if os.path.exists(filepath):
                os.remove(filepath)
            raise e
            
    except Exception as e:
        current_app.logger.error(f"Erro no upload: {e}")
        return jsonify({'sucesso': False, 'erro': 'Erro no processamento do arquivo'})


@analise_legal_bert_bp.route('/resultado/<int:doc_id>')
@login_required
def resultado(doc_id):
    """Página de resultado da análise"""
    
    documento = AnaliseLegalDocumentoBert.query.get_or_404(doc_id)
    
    # Verificar se o usuário tem acesso
    if documento.usuario_id != current_user.id:
        flash('Acesso negado', 'error')
        return redirect(url_for('analise_legal_bert.index'))
    
    # Buscar entidades e classificações relacionadas
    entidades = EntidadeExtraidaBert.query.filter_by(documento_id=doc_id).all()
    classificacoes = ClassificacaoDocumentoBert.query.filter_by(documento_id=doc_id).all()
    
    return render_template('analise_legal_bert/resultado.html',
                         documento=documento,
                         entidades=entidades,
                         classificacoes=classificacoes)


@analise_legal_bert_bp.route('/historico')
@login_required
def historico():
    """Página de histórico de análises"""
    
    page = request.args.get('page', 1, type=int)
    per_page = 10
    
    documentos = AnaliseLegalDocumentoBert.query.filter_by(
        usuario_id=current_user.id
    ).order_by(
        AnaliseLegalDocumentoBert.created_at.desc()
    ).paginate(
        page=page, 
        per_page=per_page, 
        error_out=False
    )
    
    return render_template('analise_legal_bert/historico.html', documentos=documentos)


@analise_legal_bert_bp.route('/api/modelo/info')
@login_required
def api_modelo_info():
    """API para informações do modelo"""
    
    try:
        info = get_analyzer().get_model_info()
        return jsonify({'sucesso': True, 'dados': info})
    except Exception as e:
        current_app.logger.error(f"Erro ao obter info do modelo: {e}")
        return jsonify({'sucesso': False, 'erro': str(e)})


@analise_legal_bert_bp.route('/api/analise/status/<int:doc_id>')
@login_required
def api_status_analise(doc_id):
    """API para verificar status da análise"""
    
    documento = AnaliseLegalDocumentoBert.query.get_or_404(doc_id)
    
    if documento.usuario_id != current_user.id:
        return jsonify({'sucesso': False, 'erro': 'Acesso negado'})
    
    return jsonify({
        'sucesso': True,
        'status': documento.status_analise,
        'progresso': calcular_progresso_analise(documento.status_analise),
        'tempo_processamento': documento.tempo_processamento
    })


@analise_legal_bert_bp.route('/api/documento/<int:doc_id>/entidades')
@login_required
def api_entidades_documento(doc_id):
    """API para listar entidades de um documento"""
    
    documento = AnaliseLegalDocumentoBert.query.get_or_404(doc_id)
    
    if documento.usuario_id != current_user.id:
        return jsonify({'sucesso': False, 'erro': 'Acesso negado'})
    
    entidades = EntidadeExtraidaBert.query.filter_by(documento_id=doc_id).all()
    
    return jsonify({
        'sucesso': True,
        'entidades': [entidade.to_dict() for entidade in entidades]
    })


@analise_legal_bert_bp.route('/extrair-texto', methods=['POST'])
@login_required
def extrair_texto_arquivo():
    """Extrair texto de arquivo para preenchimento da área de conteúdo"""
    
    try:
        if 'arquivo' not in request.files:
            return jsonify({'sucesso': False, 'erro': 'Nenhum arquivo enviado'})
        
        arquivo = request.files['arquivo']
        
        if arquivo.filename == '':
            return jsonify({'sucesso': False, 'erro': 'Nenhum arquivo selecionado'})
        
        # Verificar extensão
        extensoes_permitidas = {'txt', 'pdf', 'docx', 'doc'}
        extensao = arquivo.filename.rsplit('.', 1)[1].lower() if arquivo.filename and '.' in arquivo.filename else ''
        
        if extensao not in extensoes_permitidas:
            return jsonify({'sucesso': False, 'erro': 'Formato de arquivo não suportado'})
        
        # Verificar tamanho (máximo 25MB)
        if arquivo.content_length and arquivo.content_length > 25 * 1024 * 1024:
            return jsonify({'sucesso': False, 'erro': 'Arquivo muito grande (máximo 25MB)'})
        
        # Salvar temporariamente
        import tempfile
        import os
        
        with tempfile.NamedTemporaryFile(delete=False, suffix=f'.{extensao}') as temp_file:
            arquivo.save(temp_file.name)
            temp_filepath = temp_file.name
        
        try:
            # Extrair texto baseado na extensão
            if extensao == 'txt':
                with open(temp_filepath, 'r', encoding='utf-8') as f:
                    texto_extraido = f.read()
            elif extensao == 'pdf':
                texto_extraido = extrair_texto_pdf(temp_filepath)
            elif extensao in ['docx', 'doc']:
                texto_extraido = extrair_texto_docx(temp_filepath)
            else:
                raise ValueError("Formato não suportado")
            
            # Limpar arquivo temporário
            os.unlink(temp_filepath)
            
            # Verificar se o texto foi extraído
            if not texto_extraido or len(texto_extraido.strip()) < 10:
                return jsonify({'sucesso': False, 'erro': 'Não foi possível extrair texto do arquivo'})
            
            current_app.logger.info(f"✅ Texto extraído de arquivo {arquivo.filename} - {len(texto_extraido)} caracteres")
            
            return jsonify({
                'sucesso': True,
                'texto': texto_extraido,
                'nome_arquivo': arquivo.filename,
                'caracteres': len(texto_extraido)
            })
            
        except Exception as e:
            # Limpar arquivo temporário em caso de erro
            if os.path.exists(temp_filepath):
                os.unlink(temp_filepath)
            raise e
            
    except Exception as e:
        current_app.logger.error(f"❌ Erro ao extrair texto: {e}")
        return jsonify({'sucesso': False, 'erro': 'Erro interno ao processar arquivo'})


@analise_legal_bert_bp.route('/excluir/<int:doc_id>', methods=['POST', 'DELETE'])
@login_required
def excluir_documento(doc_id):
    """Excluir documento de análise"""
    
    try:
        documento = AnaliseLegalDocumentoBert.query.get_or_404(doc_id)
        
        # Verificar se o usuário tem acesso
        if documento.usuario_id != current_user.id:
            flash('Acesso negado', 'error')
            return redirect(url_for('analise_legal_bert.historico'))
        
        # Excluir documento (cascade vai excluir entidades e classificações)
        db.session.delete(documento)
        db.session.commit()
        
        current_app.logger.info(f"✅ Documento {doc_id} excluído por usuário {current_user.id}")
        
        if request.method == 'DELETE' or request.is_json:
            return jsonify({'sucesso': True, 'mensagem': 'Documento excluído com sucesso'})
        else:
            flash('Documento excluído com sucesso', 'success')
            return redirect(url_for('analise_legal_bert.historico'))
            
    except Exception as e:
        current_app.logger.error(f"❌ Erro ao excluir documento {doc_id}: {e}")
        
        if request.method == 'DELETE' or request.is_json:
            return jsonify({'sucesso': False, 'erro': 'Erro interno do servidor'})
        else:
            flash('Erro ao excluir documento', 'error')
            return redirect(url_for('analise_legal_bert.historico'))


def processar_analise_documento(doc_id, titulo, conteudo, config_analise=None):
    """Processa análise de documento usando o LegalAnalyzer"""
    
    try:
        # Obter analisador
        legal_analyzer = get_analyzer()
        
        # Executar análise com configurações
        resultado_analise = legal_analyzer.analyze_document(conteudo, titulo, config_analise)
        
        # Atualizar documento no banco
        documento = AnaliseLegalDocumentoBert.query.get(doc_id)
        if documento:
            documento.status_analise = resultado_analise.get('status', 'erro')
            documento.resumo_automatico = resultado_analise.get('resumo_automatico', '')
            documento.score_confianca = resultado_analise.get('score_confianca', 0.0)
            documento.tempo_processamento = resultado_analise.get('tempo_processamento', 0.0)
            documento.modelo_utilizado = resultado_analise.get('modelo_utilizado', 'legal_analyzer')
            documento.resultado_analise = resultado_analise
            documento.tipo_documento = resultado_analise.get('classificacao', {}).get('categoria', 'indefinido')
        
        # Salvar entidades extraídas
        entidades = resultado_analise.get('entidades', [])
        for entidade_data in entidades:
            entidade = EntidadeExtraidaBert()
            entidade.documento_id = doc_id
            entidade.tipo_entidade = entidade_data.get('tipo', '')
            entidade.texto_entidade = entidade_data.get('texto', '')
            entidade.posicao_inicio = entidade_data.get('posicao_inicio', 0)
            entidade.posicao_fim = entidade_data.get('posicao_fim', 0)
            entidade.confianca = entidade_data.get('confianca', 0.0)
            entidade.contexto = entidade_data.get('contexto', '')
            entidade.normalizada = entidade_data.get('normalizada', '')
            db.session.add(entidade)
        
        # Salvar classificações
        classificacao = resultado_analise.get('classificacao', {})
        if classificacao.get('categoria'):
            classificacao_obj = ClassificacaoDocumentoBert()
            classificacao_obj.documento_id = doc_id
            classificacao_obj.categoria = classificacao.get('categoria', '')
            classificacao_obj.subcategoria = classificacao.get('subcategoria', '')
            classificacao_obj.confianca = classificacao.get('confianca', 0.0)
            classificacao_obj.justificativa = f"Classificado baseado em: {', '.join(classificacao.get('palavras_chave', []))}"
            classificacao_obj.palavras_chave = classificacao.get('palavras_chave', [])
            db.session.add(classificacao_obj)
        
        db.session.commit()
        
        current_app.logger.info(f"✅ Análise concluída para documento {doc_id}")
        return {'sucesso': True, 'resultado': resultado_analise}
        
    except Exception as e:
        current_app.logger.error(f"❌ Erro na análise do documento {doc_id}: {e}")
        
        # Atualizar status de erro
        documento = AnaliseLegalDocumentoBert.query.get(doc_id)
        if documento:
            documento.status_analise = 'erro'
            db.session.commit()
        
        return {'sucesso': False, 'erro': str(e)}


def extrair_texto_pdf(filepath):
    """Extrai texto de arquivo PDF"""
    try:
        import PyPDF2
        with open(filepath, 'rb') as file:
            reader = PyPDF2.PdfReader(file)
            text = ""
            for page in reader.pages:
                text += page.extract_text() + "\n"
        return text
    except ImportError:
        raise ValueError("PyPDF2 não instalado para processar PDFs")
    except Exception as e:
        raise ValueError(f"Erro ao processar PDF: {e}")


def extrair_texto_docx(filepath):
    """Extrai texto de arquivo DOCX"""
    try:
        import docx
        doc = docx.Document(filepath)
        text = ""
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
        return text
    except ImportError:
        raise ValueError("python-docx não instalado para processar DOCX")
    except Exception as e:
        raise ValueError(f"Erro ao processar DOCX: {e}")


def calcular_progresso_analise(status):
    """Calcula progresso baseado no status"""
    progressos = {
        'pendente': 0,
        'processando': 50,
        'concluida': 100,
        'erro': 0
    }
    return progressos.get(status, 0)


@analise_legal_bert_bp.route('/exportar/docx/<int:doc_id>')
@login_required
def exportar_docx(doc_id):
    """Exporta relatório em formato DOCX formatado"""
    
    documento = AnaliseLegalDocumentoBert.query.get_or_404(doc_id)
    
    # Verificar se o usuário tem acesso
    if documento.usuario_id != current_user.id:
        flash('Acesso negado', 'error')
        return redirect(url_for('analise_legal_bert.index'))
    
    # Buscar entidades e classificações relacionadas
    entidades = EntidadeExtraidaBert.query.filter_by(documento_id=doc_id).all()
    classificacoes = ClassificacaoDocumentoBert.query.filter_by(documento_id=doc_id).all()
    
    try:
        from docx import Document
        from docx.shared import Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.style import WD_STYLE_TYPE
        from datetime import datetime
        import tempfile
        import os
        
        # Criar novo documento
        doc = Document()
        
        # Cabeçalho
        header = doc.sections[0].header
        header_para = header.paragraphs[0]
        header_para.text = "Relatório de Análise Legal - BERTimbau"
        header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Título principal
        title = doc.add_heading('Relatório de Análise Jurídica', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Informações do documento
        doc.add_heading('Informações do Documento', level=1)
        info_table = doc.add_table(rows=0, cols=2)
        info_table.style = 'Table Grid'
        
        # Adicionar informações
        dados_info = [
            ('Título:', documento.titulo),
            ('Data de Análise:', documento.created_at.strftime('%d/%m/%Y %H:%M')),
            ('Modelo Utilizado:', documento.modelo_utilizado or 'Legal-BERTimbau'),
            ('Status:', documento.status_analise.title()),
            ('Score de Confiança:', f"{(documento.score_confianca * 100):.1f}%" if documento.score_confianca else 'N/A'),
            ('Tempo de Processamento:', f"{documento.tempo_processamento:.2f}s" if documento.tempo_processamento else 'N/A'),
            ('Tipo de Documento:', documento.tipo_documento.title() if documento.tipo_documento else 'Não classificado')
        ]
        
        for label, valor in dados_info:
            row_cells = info_table.add_row().cells
            row_cells[0].text = label
            row_cells[1].text = str(valor)
        
        # Resumo automático
        if documento.resumo_automatico:
            doc.add_heading('Resumo Executivo', level=1)
            resumo_para = doc.add_paragraph(documento.resumo_automatico)
            resumo_para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        
        # Classificações
        if classificacoes:
            doc.add_heading('Classificação do Documento', level=1)
            for classificacao in classificacoes:
                doc.add_paragraph(f"Categoria: {classificacao.categoria}", style='List Bullet')
                if classificacao.subcategoria:
                    doc.add_paragraph(f"Subcategoria: {classificacao.subcategoria}", style='List Bullet 2')
                doc.add_paragraph(f"Confiança: {(classificacao.confianca * 100):.1f}%", style='List Bullet 2')
                if classificacao.justificativa:
                    doc.add_paragraph(f"Justificativa: {classificacao.justificativa}", style='List Bullet 2')
        
        # Entidades extraídas
        if entidades:
            doc.add_heading('Entidades Extraídas', level=1)
            
            # Agrupar entidades por tipo
            entidades_by_type = {}
            for entidade in entidades:
                tipo = entidade.tipo_entidade
                if tipo not in entidades_by_type:
                    entidades_by_type[tipo] = []
                entidades_by_type[tipo].append(entidade)
            
            for tipo, lista_entidades in entidades_by_type.items():
                doc.add_heading(f'{tipo.replace("_", " ").title()} ({len(lista_entidades)} encontradas)', level=2)
                
                # Criar tabela para entidades
                ent_table = doc.add_table(rows=1, cols=4)
                ent_table.style = 'Table Grid'
                hdr_cells = ent_table.rows[0].cells
                hdr_cells[0].text = 'Texto'
                hdr_cells[1].text = 'Normalizada'
                hdr_cells[2].text = 'Confiança'
                hdr_cells[3].text = 'Contexto'
                
                for entidade in lista_entidades:
                    row_cells = ent_table.add_row().cells
                    row_cells[0].text = entidade.texto_entidade
                    row_cells[1].text = entidade.normalizada or 'N/A'
                    row_cells[2].text = f"{(entidade.confianca * 100):.1f}%"
                    row_cells[3].text = entidade.contexto[:100] + '...' if len(entidade.contexto) > 100 else entidade.contexto
        
        # Conteúdo original (limitado)
        doc.add_heading('Conteúdo Analisado (Resumo)', level=1)
        conteudo_limitado = documento.conteudo_original[:2000] + '...' if len(documento.conteudo_original) > 2000 else documento.conteudo_original
        conteudo_para = doc.add_paragraph(conteudo_limitado)
        conteudo_para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        
        # Rodapé
        footer = doc.sections[0].footer
        footer_para = footer.paragraphs[0]
        footer_para.text = f"Relatório gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')} - Sistema Legal BERTimbau"
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Salvar em arquivo temporário
        temp_dir = tempfile.gettempdir()
        filename = f"relatorio_analise_{doc_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
        filepath = os.path.join(temp_dir, filename)
        doc.save(filepath)
        
        # Enviar arquivo
        return send_file(
            filepath,
            as_attachment=True,
            download_name=filename,
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        
    except ImportError:
        flash('Biblioteca python-docx não disponível para exportação DOCX', 'error')
        return redirect(url_for('analise_legal_bert.resultado', doc_id=doc_id))
    except Exception as e:
        current_app.logger.error(f"Erro ao gerar DOCX: {e}")
        flash('Erro ao gerar relatório DOCX', 'error')
        return redirect(url_for('analise_legal_bert.resultado', doc_id=doc_id))


@analise_legal_bert_bp.route('/exportar/json/<int:doc_id>')
@login_required
def exportar_json(doc_id):
    """Exporta dados em formato JSON com campos extraídos"""
    
    documento = AnaliseLegalDocumentoBert.query.get_or_404(doc_id)
    
    # Verificar se o usuário tem acesso
    if documento.usuario_id != current_user.id:
        return jsonify({'erro': 'Acesso negado'}), 403
    
    # Buscar entidades e classificações relacionadas
    entidades = EntidadeExtraidaBert.query.filter_by(documento_id=doc_id).all()
    classificacoes = ClassificacaoDocumentoBert.query.filter_by(documento_id=doc_id).all()
    
    try:
        # Estruturar dados para JSON
        dados_json = {
            'documento': {
                'id': documento.id,
                'titulo': documento.titulo,
                'data_analise': documento.created_at.isoformat(),
                'status_analise': documento.status_analise,
                'modelo_utilizado': documento.modelo_utilizado,
                'tipo_documento': documento.tipo_documento,
                'score_confianca': documento.score_confianca,
                'tempo_processamento': documento.tempo_processamento,
                'resumo_automatico': documento.resumo_automatico,
                'conteudo_original': documento.conteudo_original
            },
            'entidades_extraidas': [],
            'classificacoes': [],
            'estatisticas': {
                'total_entidades': len(entidades),
                'tipos_entidades': len(set(e.tipo_entidade for e in entidades)),
                'total_classificacoes': len(classificacoes)
            },
            'metadata': {
                'versao_exportacao': '1.0',
                'data_exportacao': datetime.now().isoformat(),
                'usuario_id': current_user.id,
                'sistema': 'Legal-BERTimbau'
            }
        }
        
        # Adicionar entidades
        entidades_by_type = {}
        for entidade in entidades:
            entidade_data = {
                'id': entidade.id,
                'tipo_entidade': entidade.tipo_entidade,
                'texto_entidade': entidade.texto_entidade,
                'normalizada': entidade.normalizada,
                'posicao_inicio': entidade.posicao_inicio,
                'posicao_fim': entidade.posicao_fim,
                'confianca': entidade.confianca,
                'contexto': entidade.contexto,
                'data_extracao': entidade.created_at.isoformat() if entidade.created_at else None
            }
            dados_json['entidades_extraidas'].append(entidade_data)
            
            # Agrupar por tipo para estatísticas
            tipo = entidade.tipo_entidade
            if tipo not in entidades_by_type:
                entidades_by_type[tipo] = []
            entidades_by_type[tipo].append(entidade_data)
        
        # Adicionar classificações
        for classificacao in classificacoes:
            classificacao_data = {
                'id': classificacao.id,
                'categoria': classificacao.categoria,
                'subcategoria': classificacao.subcategoria,
                'confianca': classificacao.confianca,
                'justificativa': classificacao.justificativa,
                'palavras_chave': classificacao.palavras_chave,
                'data_classificacao': classificacao.created_at.isoformat() if classificacao.created_at else None
            }
            dados_json['classificacoes'].append(classificacao_data)
        
        # Adicionar estatísticas detalhadas por tipo
        dados_json['estatisticas']['entidades_por_tipo'] = {
            tipo: len(lista) for tipo, lista in entidades_by_type.items()
        }
        
        # Adicionar análise de confiança
        if entidades:
            confidencias = [e.confianca for e in entidades]
            dados_json['estatisticas']['confianca_media'] = sum(confidencias) / len(confidencias)
            dados_json['estatisticas']['confianca_minima'] = min(confidencias)
            dados_json['estatisticas']['confianca_maxima'] = max(confidencias)
        
        # Preparar resposta
        filename = f"analise_legal_{doc_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
        
        response = jsonify(dados_json)
        response.headers['Content-Disposition'] = f'attachment; filename={filename}'
        response.headers['Content-Type'] = 'application/json; charset=utf-8'
        
        return response
        
    except Exception as e:
        current_app.logger.error(f"Erro ao gerar JSON: {e}")
        return jsonify({'erro': 'Erro ao gerar exportação JSON'}), 500